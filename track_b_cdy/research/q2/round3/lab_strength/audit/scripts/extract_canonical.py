"""Read only two specified XLSX members and emit unfit canonical records.

No workbook writes, formula evaluation, imputation or row filtering. Nominal
class membership is separate from measured-at-failure numerical values.
"""
from __future__ import annotations

from collections import Counter
import csv
import hashlib
import io
import json
import math
from pathlib import Path
import re
import sys
import zipfile

sys.dont_write_bytecode = True
import openpyxl

ROOT = Path(__file__).resolve().parents[7]
AUDIT = ROOT / "track_b/research/q2/round3/lab_strength/audit"
RESULTS = AUDIT / "results"
ZIP = ROOT / "track_b/research/q2/round3/sources/strength/mamot2018_supplement.zip"
ZIP_SHA = "5fcd735a88654af1f4c150313ee79afb855a0ae67ed9319a851d8e125fa420a0"
MEMBERS = {
    "experiments": ("Supplement/data_experiments.xlsx", "ed9517ae4889192bd1a3278c68b4414c391126dbe9e3133336923795fbbbb20b"),
    "coefficients": ("Supplement/data_frictioncoeff_cohesion_Fig. 8.xlsx", "34ec81f6adb579cd16e8b39a1934d5b786a633f96f12d2a05f65ed7b2fc1a165"),
}
FIELDS = [
    ("experiment_id", "A", "identifier"), ("specimen_label", "B", "identifier"),
    ("ice_layer_thickness_mm", "C", "mm"), ("temperature_target_c", "D", "degC"),
    ("temperature_at_failure_c", "E", "degC"), ("normal_stress_target_kpa", "F", "kPa"),
    ("normal_stress_at_failure_kpa", "G", "kPa"), ("deformation_start_s", "H", "s"),
    ("shear_start_to_end_s", "I", "s"), ("rupture_time_s", "J", "s"),
    ("AE_onset_s", "K", "s"), ("peak_shear_kpa", "L", "kPa"),
    ("failure_type_code", "M", "code"), ("failure_type_description", "N", "category"),
    ("shear_rate_mm_min", "O", "mm/min"), ("shear_deformation_mm", "P", "mm"),
]


def sha(raw): return hashlib.sha256(raw).hexdigest()


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False)+"\n")


def write_csv(path, rows, fields=None):
    path.parent.mkdir(parents=True, exist_ok=True)
    fields = fields or list(rows[0])
    with path.open("w", newline="") as f:
        w=csv.DictWriter(f, fieldnames=fields);w.writeheader();w.writerows(rows)


def read_inputs():
    raw=ZIP.read_bytes();assert sha(raw)==ZIP_SHA
    records={}
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        for key,(member, expected) in MEMBERS.items():
            data=z.read(member);assert sha(data)==expected
            records[key] = {"member":member,"bytes":len(data),"sha256":expected,
                            "crc32":f"{z.getinfo(member).CRC:08x}", "raw":data}
    return records


def number(value, reference, issues):
    if value is None or value == "":
        issues.append({"reference":reference,"kind":"missing_numeric","raw_value":value})
        return None
    if isinstance(value, bool):
        issues.append({"reference":reference,"kind":"unexpected_boolean","raw_value":value})
        return None
    if isinstance(value,(int,float)):
        if math.isfinite(value): return float(value)
        issues.append({"reference":reference,"kind":"nonfinite_numeric","raw_value":str(value)})
        return None
    if isinstance(value,str) and re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?",value.strip()):
        issues.append({"reference":reference,"kind":"numeric_string_parsed_without_imputation","raw_value":value})
        return float(value)
    issues.append({"reference":reference,"kind":"unparseable_numeric","raw_value":str(value)})
    return None


def extract():
    assert str(ROOT)=="/Users/maxchen/Documents/MathorCup/worktrees/track-b-q2"
    inputs=read_inputs(); issues=[]
    books={k:openpyxl.load_workbook(io.BytesIO(v["raw"]),data_only=False,read_only=False) for k,v in inputs.items()}
    cached={k:openpyxl.load_workbook(io.BytesIO(v["raw"]),data_only=True,read_only=False) for k,v in inputs.items()}
    s=books["experiments"]["Sheet1"];cs=cached["experiments"]["Sheet1"]
    assert s.max_row==142 and s.max_column==16
    records=[]
    for row in range(2,143):
        raw={name:s[f"{col}{row}"].value for name,col,_ in FIELDS}
        converted={}
        field_sources={}
        row_issues=[]
        for name,col,unit in FIELDS:
            cell=s[f"{col}{row}"]; reference=f"{inputs['experiments']['member']}#Sheet1!{cell.coordinate}"
            value=cs[cell.coordinate].value if cell.data_type=="f" else cell.value
            field_sources[name]={"cell":cell.coordinate,"header_cell":f"{col}1","source_header":s[f"{col}1"].value,"unit":unit,"source_cell_type":cell.data_type}
            converted[name]=value if name in ["experiment_id","specimen_label","failure_type_description"] else number(value,reference,row_issues)
        obsid=f"mamot2018-test-r{row:03d}"
        group=(converted["experiment_id"],converted["specimen_label"])
        complete=all(converted[name] is not None for name,_,_ in FIELDS)
        t,tn=converted["temperature_at_failure_c"],converted["temperature_target_c"]
        p,pn=converted["normal_stress_at_failure_kpa"],converted["normal_stress_target_kpa"]
        low=pn is not None and 100<=pn<=400
        warm=tn is not None and -4<=tn<=-.5
        flags=[i["kind"]+":"+i["reference"].split("!")[-1] for i in row_issues]
        record={"observation_id":obsid,"source_zip_path":ZIP.relative_to(ROOT).as_posix(),"source_zip_sha256":ZIP_SHA,
                "workbook_member":inputs["experiments"]["member"],"workbook_member_sha256":inputs["experiments"]["sha256"],
                "sheet":"Sheet1","source_row":row,"test_key_experiment_and_specimen":json.dumps(group),**converted,
                "physical_specimen_id":None,"specimen_label_cluster_candidate":converted["specimen_label"],
                "specimen_independence_status":"unknown_global_identity; recurring label is not a proven new specimen or proven reused rock-pair UID",
                "complete_source_row":complete,"in_nominal_low_stress_group":low,"in_nominal_warm_group":warm,
                "in_nominal_warm_low_group":low and warm,
                "in_nominal_equation_temperature_stress_domain":low and tn is not None and -8<=tn<=-.5,
                "in_actual_equation_temperature_stress_domain":t is not None and p is not None and -8<=t<=-.5 and 100<=p<=400,
                "normal_stress_actual_minus_target_kpa":p-pn if p is not None and pn is not None else None,
                "temperature_actual_minus_target_c":t-tn if t is not None and tn is not None else None,
                "strain_rate_per_s_derived":converted["shear_rate_mm_min"]/(60*converted["ice_layer_thickness_mm"]) if converted["shear_rate_mm_min"] is not None and converted["ice_layer_thickness_mm"] else None,
                "strain_rate_derivation":"shear_rate_mm_min/(60*ice_layer_thickness_mm); derived, not an independently measured column",
                "time_origin_status":"source headers retained; J/K absolute vs relative origins not determined by workbook alone",
                "field_cells_json":json.dumps(field_sources,ensure_ascii=False,sort_keys=True),
                "raw_typed_values_json":json.dumps(raw,ensure_ascii=False,sort_keys=True),
                "audit_flags_json":json.dumps(flags),"evaluation_stratum":"external_transfer","independent_landslide_events":0}
        records.append(record);issues.extend({**i,"observation_id":obsid} for i in row_issues)
    assert len(records)==141 and len({r["observation_id"] for r in records})==141
    write_csv(RESULTS/"observations.csv",records)

    coeff=[];s=books["coefficients"]["Sheet1"]
    blocks=[("nominal_100_400",100,400,2,range(2,11),[5,6,7,8,12,13,14,15],9),
            ("nominal_100_800",100,800,19,range(2,7),[22,23,24,25,28,29,30,31],None)]
    names=["friction_temperature_mean_c","friction_temperature_error_c","friction_coefficient_mu","friction_coefficient_error",
           "cohesion_temperature_mean_c","cohesion_temperature_error_c","cohesion_kpa","cohesion_error_kpa"]
    for block,pmin,pmax,header_row,cols,rows,angle in blocks:
        for col in cols:
            cells={name:s.cell(r,col) for name,r in zip(names,rows)}
            vals={name:number(cell.value,f"{inputs['coefficients']['member']}#Sheet1!{cell.coordinate}",issues) for name,cell in cells.items()}
            target=number(s.cell(header_row,col).value,f"{inputs['coefficients']['member']}#Sheet1!{s.cell(header_row,col).coordinate}",issues)
            sourcecells={name:cell.coordinate for name,cell in cells.items()};sourcecells["temperature_target_c"]=s.cell(header_row,col).coordinate
            angle_cell=s.cell(angle,col) if angle is not None else None
            assert vals["friction_temperature_mean_c"]==vals["cohesion_temperature_mean_c"]
            assert vals["friction_temperature_error_c"]==vals["cohesion_temperature_error_c"]
            coeff.append({"aggregate_id":f"mamot2018-{block}-col{openpyxl.utils.get_column_letter(col)}",
                "source_zip_path":ZIP.relative_to(ROOT).as_posix(),"source_zip_sha256":ZIP_SHA,
                "workbook_member":inputs["coefficients"]["member"],"workbook_member_sha256":inputs["coefficients"]["sha256"],
                "sheet":"Sheet1","nominal_stress_group":block,"normal_stress_target_min_kpa":pmin,"normal_stress_target_max_kpa":pmax,
                "temperature_target_c":target,**vals,
                "friction_angle_degrees_source":angle_cell.value if angle_cell is not None else None,
                "friction_angle_source_status":"label_present_values_blank" if angle_cell is not None else "no_angle_row_in_block",
                "friction_angle_source_cell":angle_cell.coordinate if angle_cell is not None else None,
                "field_cells_json":json.dumps(sourcecells,sort_keys=True),
                "error_semantics":"workbook labels only say _error; do not infer SD/SE/confidence type from this file alone",
                "observation_role":"author_derived_temperature_stratum_summary; not an additional independent test",
                "evaluation_stratum":"external_transfer","independent_landslide_events":0})
    assert len(coeff)==14
    write_csv(RESULTS/"coefficient_strata.csv",coeff)
    write_json(RESULTS/"normalization_issues.json",issues)
    schema={"schema_version":"mamot2018-lab-audit-v1","no_models_fitted":True,
        "observations":{"file":"observations.csv","rows":141,"one_row":"one reported shear test, not a proven independent physical specimen",
            "raw_column_map":{name:{"column":col,"unit":unit} for name,col,unit in FIELDS},
            "nominal_membership":"Use temperature_target_c and normal_stress_target_kpa; never filter stress classes by measured normal_stress_at_failure_kpa",
            "missingness_policy":"all source rows retained; missing or unparseable numeric values remain null and issue records; no zero fill",
            "numeric_strings":"F135 and G135 source strings parse exactly to200 and200.4; source raw/type/cell retained",
            "grouping":{"reported_tests":141,"experiment_conditions":len(set(r['experiment_id'] for r in records)),
                "specimen_label_counts":dict(Counter(r['specimen_label'] for r in records)),
                "physical_specimen_uid":"not provided; do not infer unique independent specimen count",
                "specimen_cluster_caution":"spec labels recur across conditions; cluster as sensitivity assumption only until physical rock-pair mapping is established"}},
        "coefficient_strata":{"file":"coefficient_strata.csv","rows":14,"role":"author-derived summaries, not additional141 observations",
            "blocks":{"nominal_100_400":9,"nominal_100_800":5}},
        "group_counts":{"nominal_low_stress":sum(r['in_nominal_low_stress_group'] for r in records),
                        "nominal_warm":sum(r['in_nominal_warm_group'] for r in records),
                        "nominal_warm_low":sum(r['in_nominal_warm_low_group'] for r in records)},
        "input_record":{"source_zip":dict(path=ZIP.relative_to(ROOT).as_posix(),bytes=ZIP.stat().st_size,sha256=ZIP_SHA),
            "workbook_members":[{k:v for k,v in item.items() if k!='raw'} for item in inputs.values()]},
        "software":{"python":sys.version,"openpyxl":openpyxl.__version__},
        "command":"track_b/.venv/bin/python track_b/research/q2/round3/lab_strength/audit/scripts/extract_canonical.py"}
    assert schema['group_counts']=={'nominal_low_stress':121,'nominal_warm':83,'nominal_warm_low':63}
    write_json(RESULTS/"schema.json",schema)
    print(json.dumps({"observations":141,"coefficient_strata":14,"groups":schema['group_counts'],"issues":len(issues),"observations_sha256":sha((RESULTS/'observations.csv').read_bytes())}))


if __name__=="__main__":extract()
