"""Full read-only OOXML/openpyxl audit of the two authorized lab workbooks.

Audits every coordinate in every declared sheet rectangle, all underlying XML
cells/strings/parts, labels, cached values, formulas, comments and repeated IDs.
Only arithmetic/descriptive reconciliation is performed; no models are fitted.
"""
from __future__ import annotations

from collections import Counter, defaultdict
import csv
import hashlib
import io
import json
from pathlib import Path, PurePosixPath
import posixpath
import statistics
import sys
import xml.etree.ElementTree as ET
import zipfile

sys.dont_write_bytecode = True
import openpyxl
from openpyxl.utils import get_column_letter, range_boundaries
from extract_canonical import ROOT, AUDIT, RESULTS, ZIP, ZIP_SHA, MEMBERS, FIELDS, read_inputs, sha, write_csv, write_json

NS = {"m":"http://schemas.openxmlformats.org/spreadsheetml/2006/main",
      "r":"http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "p":"http://schemas.openxmlformats.org/package/2006/relationships"}


def json_string(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False, default=str)


def node_text(node):
    return "".join(node.itertext()) if node is not None else None


def load_csv(path):
    with path.open(newline="") as f: return list(csv.DictReader(f))


def as_float(value):
    return float(value) if value is not None and value!="" else None


def main():
    assert str(ROOT)=="/Users/maxchen/Documents/MathorCup/worktrees/track-b-q2"
    inputs=read_inputs()
    before={name:sha((RESULTS/name).read_bytes()) for name in ["observations.csv","coefficient_strata.csv","schema.json","normalization_issues.json"]}
    parts=[]; cells=[]; strings=[]; labels=[]; sheet_records=[]; workbook_records=[]; failures=[]
    raw_inventory_coverage=[]
    for key,item in inputs.items():
        archive=zipfile.ZipFile(io.BytesIO(item["raw"]))
        book=openpyxl.load_workbook(io.BytesIO(item["raw"]),read_only=False,data_only=False)
        cached=openpyxl.load_workbook(io.BytesIO(item["raw"]),read_only=False,data_only=True)
        for part in archive.infolist():
            b=archive.read(part)
            parts.append(dict(workbook_key=key,workbook_member=item["member"],part=part.filename,
                              bytes=len(b),sha256=sha(b),crc32=f"{part.CRC:08x}"))
        wb_xml=ET.fromstring(archive.read("xl/workbook.xml"))
        rel_xml=ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        rels={n.attrib["Id"]:n.attrib for n in rel_xml}
        external=[]
        for part in archive.namelist():
            if part.endswith(".rels"):
                tree=ET.fromstring(archive.read(part))
                external.extend({"part":part,**n.attrib} for n in tree if n.attrib.get("TargetMode")=="External")
        if "xl/sharedStrings.xml" in archive.namelist():
            st=ET.fromstring(archive.read("xl/sharedStrings.xml"))
            for index,node in enumerate(st.findall("m:si",NS)):
                text="".join(n.text or "" for n in node.iter("{"+NS["m"]+"}t"))
                strings.append(dict(workbook_key=key,workbook_member=item["member"],shared_string_index=index,
                    decoded_text=text,xml_serialization=ET.tostring(node,encoding="unicode")))
        sheet_paths={}
        for sh in wb_xml.findall("m:sheets/m:sheet",NS):
            rel=rels[sh.attrib["{"+NS["r"]+"}id"]]
            target=rel["Target"]
            path=target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join("xl",target))
            sheet_paths[sh.attrib["name"]]=(path,sh.attrib)
        for ws in book:
            path,sh_attributes=sheet_paths[ws.title]
            raw=archive.read(path); tree=ET.fromstring(raw)
            outxml=RESULTS/"worksheet_xml"/f"{key}_{PurePosixPath(path).name}"
            outxml.parent.mkdir(parents=True,exist_ok=True);outxml.write_bytes(raw)
            xmlcells={n.attrib["r"]:n for n in tree.findall("m:sheetData/m:row/m:c",NS)}
            row_attrs={int(n.attrib["r"]):n.attrib for n in tree.findall("m:sheetData/m:row",NS)}
            column_attrs=[n.attrib for n in tree.findall("m:cols/m:col",NS)]
            formulas=0;errors=0;comments=0;nonempty=0;cached_diffs=0;explicit_blanks=0
            max_nonempty_row=max_nonempty_col=0
            for row in range(1,ws.max_row+1):
                for col in range(1,ws.max_column+1):
                    coordinate=f"{get_column_letter(col)}{row}"; cell=ws[coordinate];cache=cached[ws.title][coordinate]
                    node=xmlcells.get(coordinate); f=node.find("m:f",NS) if node is not None else None
                    v=node.find("m:v",NS) if node is not None else None
                    inline=node.find("m:is",NS) if node is not None else None
                    formula=f.text if f is not None else None
                    has_formula=f is not None or cell.data_type=="f"
                    value=cell.value; cached_value=cache.value
                    is_blank=value is None and not has_formula
                    formulas+=int(has_formula);errors+=int(cell.data_type=="e" or cache.data_type=="e")
                    comments+=int(cell.comment is not None);nonempty+=int(value is not None)
                    if value is not None:max_nonempty_row=max(max_nonempty_row,row);max_nonempty_col=max(max_nonempty_col,col)
                    if not has_formula and value!=cached_value:
                        cached_diffs+=1;failures.append(dict(kind="nonformula_cached_value_mismatch",workbook=key,sheet=ws.title,cell=coordinate))
                    if node is not None and is_blank:explicit_blanks+=1
                    if is_blank:role="blank_not_invented_measurement"
                    elif key=="experiments" and row==1:role="source_header"
                    elif key=="experiments":role="reported_test_data"
                    elif col==1:role="source_label_or_note"
                    elif row in (2,19):role="temperature_target_header"
                    else:role="author_derived_aggregate_value"
                    hidden_col=any(int(a["min"])<=col<=int(a["max"]) and a.get("hidden")=="1" for a in column_attrs)
                    merged=[str(rng) for rng in ws.merged_cells.ranges if coordinate in rng]
                    record=dict(workbook_key=key,source_zip_path=ZIP.relative_to(ROOT).as_posix(),source_zip_sha256=ZIP_SHA,
                        workbook_member=item["member"],workbook_member_sha256=item["sha256"],sheet=ws.title,
                        worksheet_part=path,sheet_state=ws.sheet_state,cell=coordinate,row=row,column=col,column_letter=get_column_letter(col),
                        cell_role=role,xml_cell_present=node is not None,
                        xml_cell_attributes_json=json_string(node.attrib if node is not None else {}),
                        xml_type=node.attrib.get("t","n") if node is not None else None,
                        xml_value_text=v.text if v is not None else None,
                        xml_formula_text=formula,xml_formula_attributes_json=json_string(f.attrib if f is not None else {}),
                        xml_inline_text=node_text(inline),decoded_value_json=json_string(value),
                        cached_value_json=json_string(cached_value),openpyxl_type=cell.data_type,cached_type=cache.data_type,
                        has_formula=has_formula,is_error=cell.data_type=="e" or cache.data_type=="e",is_blank=is_blank,
                        explicit_styled_or_empty_xml_blank=node is not None and is_blank,style_id=cell.style_id,
                        number_format=cell.number_format,row_hidden=row_attrs.get(row,{}).get("hidden")=="1",
                        column_hidden=hidden_col,merged_ranges_json=json_string(merged),
                        comment_text=cell.comment.text if cell.comment else None,comment_author=cell.comment.author if cell.comment else None,
                        hyperlink_target=cell.hyperlink.target if cell.hyperlink else None)
                    cells.append(record)
                    if role in ["source_header","source_label_or_note","temperature_target_header"] or cell.comment:
                        labels.append({k:record[k] for k in ["workbook_key","workbook_member","sheet","cell","row","column","cell_role","decoded_value_json","xml_cell_attributes_json","comment_text","comment_author"]})
            logical={f"{get_column_letter(c)}{r}" for r in range(1,ws.max_row+1) for c in range(1,ws.max_column+1)}
            outside=set(xmlcells)-logical
            if outside:failures.append(dict(kind="xml_cells_outside_audit_rectangle",workbook=key,sheet=ws.title,cells=sorted(outside)))
            sr=dict(workbook_key=key,workbook_member=item["member"],sheet=ws.title,worksheet_part=path,
                    original_worksheet_xml_path=outxml.relative_to(ROOT).as_posix(),worksheet_xml_bytes=len(raw),worksheet_xml_sha256=sha(raw),
                    sheet_attributes=sh_attributes,declared_dimension=tree.find("m:dimension",NS).attrib if tree.find("m:dimension",NS) is not None else None,
                    openpyxl_dimension=ws.calculate_dimension(),max_row=ws.max_row,max_column=ws.max_column,
                    audited_coordinate_count=ws.max_row*ws.max_column,explicit_xml_cell_count=len(xmlcells),
                    nonempty_value_count=nonempty,explicit_xml_blank_count=explicit_blanks,
                    nonempty_max_row=max_nonempty_row,nonempty_max_column=max_nonempty_col,
                    formula_count=formulas,error_count=errors,comment_count=comments,nonformula_cached_mismatches=cached_diffs,
                    merged_ranges=[str(r) for r in ws.merged_cells.ranges],rows_metadata=row_attrs,columns_metadata=column_attrs,
                    sheet_properties=ET.tostring(tree.find("m:sheetPr",NS),encoding="unicode") if tree.find("m:sheetPr",NS) is not None else None,
                    auto_filter_ref=ws.auto_filter.ref,freeze_panes=str(ws.freeze_panes) if ws.freeze_panes else None,
                    table_names=list(ws.tables),conditional_formatting_rule_count=len(ws.conditional_formatting))
            sheet_records.append(sr);raw_inventory_coverage.append(dict(workbook=key,sheet=ws.title,all_xml_cells_covered=not outside))
        workbook_records.append(dict(workbook_key=key,workbook_member=item["member"],bytes=item["bytes"],sha256=item["sha256"],
            sheet_names=book.sheetnames,defined_names={name:str(value) for name,value in book.defined_names.items()},
            external_relationships=external,archive_parts=len(archive.infolist()),
            workbook_properties_xml=ET.tostring(wb_xml.find("m:workbookPr",NS),encoding="unicode") if wb_xml.find("m:workbookPr",NS) is not None else None,
            calculation_properties_xml=ET.tostring(wb_xml.find("m:calcPr",NS),encoding="unicode") if wb_xml.find("m:calcPr",NS) is not None else None,
            formula_recalculation="not performed; original cached and expression values retained; source workbooks unchanged"))

    write_csv(RESULTS/"raw_cell_inventory.csv",cells)
    write_csv(RESULTS/"source_headers_and_notes.csv",labels)
    write_csv(RESULTS/"xlsx_part_inventory.csv",parts)
    write_csv(RESULTS/"shared_string_inventory.csv",strings)
    write_json(RESULTS/"workbook_and_sheet_metadata.json",dict(workbooks=workbook_records,sheets=sheet_records))
    obs=load_csv(RESULTS/"observations.csv"); coeff=load_csv(RESULTS/"coefficient_strata.csv")
    issues=json.loads((RESULTS/"normalization_issues.json").read_text())
    raw_by_ref={(r["workbook_key"],r["sheet"],r["cell"]):r for r in cells}
    for r in obs:
        mapping=json.loads(r["field_cells_json"])
        for field,source in mapping.items():
            raw=raw_by_ref[("experiments","Sheet1",source["cell"])]
            value=json.loads(raw["decoded_value_json"])
            if field not in ["experiment_id","specimen_label","failure_type_description"]:
                assert as_float(r[field])==(None if value is None else float(value)),(r["observation_id"],field)
            else:assert r[field]==value
    summaries=[]; condition=defaultdict(list); specimens=defaultdict(list)
    for r in obs:
        condition[r["experiment_id"]].append(r);specimens[r["specimen_label"]].append(r)
    for cid,group in condition.items():
        actual_t=[as_float(r['temperature_at_failure_c']) for r in group]
        actual_s=[as_float(r['normal_stress_at_failure_kpa']) for r in group]
        target_t=sorted({as_float(r['temperature_target_c']) for r in group});target_s=sorted({as_float(r['normal_stress_target_kpa']) for r in group})
        if len(target_t)!=1 or len(target_s)!=1:failures.append(dict(kind="condition_label_has_multiple_nominal_targets",experiment_id=cid))
        summaries.append(dict(experiment_id=cid,n_reported_tests=len(group),temperature_target_c_json=json_string(target_t),
            normal_stress_target_kpa_json=json_string(target_s),specimen_labels_json=json_string([r['specimen_label'] for r in group]),
            observation_ids_json=json_string([r['observation_id'] for r in group]),source_rows_json=json_string([int(r['source_row']) for r in group]),
            actual_temperature_min_c=min(actual_t),actual_temperature_max_c=max(actual_t),actual_temperature_mean_c=statistics.mean(actual_t),
            actual_normal_stress_min_kpa=min(actual_s),actual_normal_stress_max_kpa=max(actual_s),actual_normal_stress_mean_kpa=statistics.mean(actual_s),
            independence_status="reported repetitions within nominal condition; independent ice specimen/rock-pair identities unverified"))
    write_csv(RESULTS/"condition_replicates.csv",summaries)
    specimen_rows=[]
    for label,group in sorted(specimens.items()):
        specimen_rows.append(dict(specimen_label=label,n_reported_tests=len(group),n_experiment_conditions=len(set(r['experiment_id'] for r in group)),
            n_nominal_temperature_classes=len(set(r['temperature_target_c'] for r in group)),n_nominal_stress_classes=len(set(r['normal_stress_target_kpa'] for r in group)),
            observation_ids_json=json_string([r['observation_id'] for r in group]),source_cells_json=json_string([f"Sheet1!B{r['source_row']}" for r in group]),
            physical_specimen_id=None,independence_status="recurring local label only; no workbook global physical UID; cluster sensitivity assumption possible but not proven rock-pair mapping"))
    write_csv(RESULTS/"specimen_label_repeats.csv",specimen_rows)
    fullrows=Counter(r['raw_typed_values_json'] for r in obs)
    keys=Counter((r['experiment_id'],r['specimen_label']) for r in obs)
    identical_continuous=defaultdict(list)
    for r in obs:identical_continuous[tuple(r[k] for k in ['temperature_at_failure_c','normal_stress_at_failure_kpa','peak_shear_kpa'])].append(r['observation_id'])
    duplicates=dict(exact_full_data_row_duplicate_groups=[dict(raw_typed_values_json=k,n=v) for k,v in fullrows.items() if v>1],
        duplicate_experiment_specimen_keys=[dict(experiment_id=k[0],specimen_label=k[1],n=v) for k,v in keys.items() if v>1],
        exact_temperature_stress_peak_duplicate_groups=[dict(values=list(k),observation_ids=v) for k,v in identical_continuous.items() if len(v)>1],
        specimen_label_count=len(specimens),specimen_label_counts={k:len(v) for k,v in specimens.items()},
        independent_physical_specimen_count=None,
        specimen_identity_note="Workbook has8labels including1spec8, but no IDs linking labels across tests to physical rock pairs. The paper's seven-rock-pair context is owned by the separate literature review; no rock-pair identity is asserted here.")
    write_json(RESULTS/"duplicates_and_independence.json",duplicates)
    reconciliations=[]
    for c in coeff:
        target=as_float(c['temperature_target_c']);low=as_float(c['normal_stress_target_min_kpa']);high=as_float(c['normal_stress_target_max_kpa'])
        group=[r for r in obs if as_float(r['temperature_target_c'])==target and low<=as_float(r['normal_stress_target_kpa'])<=high]
        vals=[as_float(r['temperature_at_failure_c']) for r in group]
        mean=statistics.mean(vals);sd=statistics.stdev(vals);psd=statistics.pstdev(vals)
        stated=as_float(c['friction_temperature_mean_c']);err=as_float(c['friction_temperature_error_c'])
        reconciliations.append(dict(aggregate_id=c['aggregate_id'],temperature_target_c=target,nominal_stress_group=c['nominal_stress_group'],
            group_membership="target nominal classes, never measured-stress cutoff",n_reported_tests=len(group),
            observation_ids_json=json_string([r['observation_id'] for r in group]),published_mean_temperature_c=stated,
            recomputed_raw_test_mean_temperature_c=mean,mean_difference_published_minus_raw=stated-mean,
            published_temperature_error_c=err,recomputed_raw_temperature_sample_sd_c=sd,
            recomputed_raw_temperature_population_sd_c=psd,error_difference_from_sample_sd=err-sd,
            published_mean_within_0_0051_rounding_tolerance=abs(stated-mean)<=.0051,
            temperature_error_matches_sample_sd_within_1e_10=abs(err-sd)<1e-10,
            interpretation="descriptive reconciliation only; errors retain original source label, no regression or model fitted"))
    write_csv(RESULTS/"coefficient_temperature_reconciliation.csv",reconciliations)
    missing=[]
    for name,_,unit in FIELDS:
        missing.append(dict(field=name,unit=unit,n_rows=141,n_missing=sum(r[name]=='' for r in obs),
            n_numeric_source_strings=sum(i['kind']=='numeric_string_parsed_without_imputation' and i['reference'].endswith('!'+json.loads(r['field_cells_json'])[name]['cell']) for r in obs for i in issues)))
    write_csv(RESULTS/"field_missingness.csv",missing)
    lab_nominal=[r for r in obs if r['in_nominal_equation_temperature_stress_domain']=='True']
    lab_actual=[r for r in obs if r['in_actual_equation_temperature_stress_domain']=='True']
    changed=[r for r in obs if r['in_nominal_equation_temperature_stress_domain']!=r['in_actual_equation_temperature_stress_domain']]
    domain=dict(grouping_for_parent_fits="nominal temperature_target_c/normal_stress_target_kpa only; actual continuous values remain response/covariate inputs",
        nominal_low_stress_n=sum(100<=as_float(r['normal_stress_target_kpa'])<=400 for r in obs),
        actual_stress_100_400_n=sum(100<=as_float(r['normal_stress_at_failure_kpa'])<=400 for r in obs),
        nominal_warm_n=sum(-4<=as_float(r['temperature_target_c'])<=-.5 for r in obs),
        nominal_warm_low_n=sum(-4<=as_float(r['temperature_target_c'])<=-.5 and 100<=as_float(r['normal_stress_target_kpa'])<=400 for r in obs),
        nominal_equation_temperature_stress_domain_n=len(lab_nominal),actual_equation_temperature_stress_domain_n=len(lab_actual),
        target_vs_actual_domain_flag_changed_n=len(changed),changed_observation_ids=[r['observation_id'] for r in changed],
        actual_normal_stress_above400_with_nominal_low_n=sum(as_float(r['normal_stress_at_failure_kpa'])>400 and 100<=as_float(r['normal_stress_target_kpa'])<=400 for r in obs),
        temperature_target_values=sorted(set(as_float(r['temperature_target_c']) for r in obs)),
        normal_stress_target_values=sorted(set(as_float(r['normal_stress_target_kpa']) for r in obs)))
    write_json(RESULTS/"nominal_actual_domain_audit.json",domain)
    failure_map=defaultdict(set)
    for r in obs:failure_map[r['failure_type_code']].add(r['failure_type_description'])
    semantic=dict(failure_type_code_to_descriptions={k:sorted(v) for k,v in failure_map.items()},
        failure_type_counts=dict(Counter(r['failure_type_description'] for r in obs)),
        all_reported_tests_retained=True,no_model_fits=True,no_source_values_overwritten=True,
        timing_note="H-J-K source labels differ; no unverified common origin assumed, no AE lead-time derived",
        equation_input_note="temperature_target and normal_stress_target are group labels, not replacements for actual-at-failure values",
        coefficient_note="14author-derived temperature strata summarize same experiments; do not append as14independent shear tests",
        friction_angle_source_note="100-400block row9 labels friction angle_degree but9numeric cells B9:J9 are blank; 100-800block has no angle row; no atan(mu) imputation written into source-angle column",
        figure_label_note="source filename/headings say Fig.8a/8b; preserved literally; no automatic remapping to published figure numbering",
        physical_specimen_note="No stable physical specimen UID inferred. Eight specimen labels are not proven eight independent rock pairs; 141rows are tests, not141independent specimens.")
    write_json(RESULTS/"semantics_and_ambiguities.json",semantic)
    assert len(cells)==4074
    assert len(obs)==141 and len(coeff)==14 and len(condition)==32 and len(specimens)==8
    assert all(r['all_xml_cells_covered'] for r in raw_inventory_coverage)
    assert len(issues)==2 and all(i['kind']=='numeric_string_parsed_without_imputation' for i in issues)
    assert not failures,failures
    assert all(sha((RESULTS/name).read_bytes())==value for name,value in before.items())
    assert sha(ZIP.read_bytes())==ZIP_SHA
    validation=dict(all_checks_passed=True,raw_cell_coordinates=4074,nonempty_cells=sum(not c['is_blank'] for c in cells),
        sheets_audited=len(sheet_records),blank_sheets=sum(s['nonempty_value_count']==0 for s in sheet_records),
        formulas=sum(s['formula_count'] for s in sheet_records),cached_formula_missing_or_stale_status="no formulas in either source; no calculation required or performed",
        source_error_cells=sum(s['error_count'] for s in sheet_records),comments=sum(s['comment_count'] for s in sheet_records),
        nonformula_cached_mismatches=sum(s['nonformula_cached_mismatches'] for s in sheet_records),
        xml_parts_indexed=len(parts),shared_strings_indexed=len(strings),source_headers_notes_cells=len(labels),
        normalized_observations=141,normalized_aggregate_records=14,numeric_string_conversions=2,
        missing_source_fields=sum(r['n_missing'] for r in missing),all_normalized_fields_trace_to_original_cells=True,
        raw_inventory_coverage=raw_inventory_coverage,all_source_hashes_unchanged=True,canonical_early_outputs_unchanged=True,
        failures=failures,independent_events=0,evaluation_stratum="external_transfer",
        statistical_fits_performed=0,physical_independent_specimens=None)
    write_json(RESULTS/"validation.json",validation)
    input_manifest=dict(source_zip_path=ZIP.relative_to(ROOT).as_posix(),source_zip_bytes=ZIP.stat().st_size,source_zip_sha256=ZIP_SHA,
        canonical_url_or_product_id="https://tc.copernicus.org/articles/12/3333/2018/tc-12-3333-2018-supplement.zip",
        product_version="Mamot et al.2018 official supplementary ZIP; two specified workbook members only",
        retrieval_status="reused existing frozen local source; no new network retrieval",
        retrieved_at_utc=None,retrieval_time_note="original acquisition remains in earlier strength source register; no new acquisition timestamp invented",
        actual_read_scope="all sheets, every declared worksheet coordinate and XML part in two specified XLSX; no other ZIP workbook opened",
        workbook_members=[{k:v for k,v in item.items() if k!='raw'} for item in inputs.values()],
        scripts=[dict(path=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=sha(p.read_bytes())) for p in sorted((AUDIT/'scripts').glob('*.py'))],
        software={"python":sys.version,"openpyxl":openpyxl.__version__,"reader_modes":"openpyxl formula and data_only views; stdlib zip/XML authoritative raw cell inventory"})
    write_json(RESULTS/"input_record.json",input_manifest)
    report=dict(schema_version="mamot2018-complete-workbook-audit-v1",canonical_schema_path=(RESULTS/'schema.json').relative_to(ROOT).as_posix(),
        summary=validation,domain=domain,duplicates=duplicates,semantics=semantic,
        coefficient_temperature_reconciliation=dict(records=14,
            published_means_within_rounding_tolerance_n=sum(r['published_mean_within_0_0051_rounding_tolerance'] for r in reconciliations),
            published_errors_matching_sample_sd_n=sum(r['temperature_error_matches_sample_sd_within_1e_10'] for r in reconciliations),
            note="Any mismatch is retained for parent literature/model review; no error type relabeled or source correction performed"),
        reproduction_commands=["track_b/.venv/bin/python track_b/research/q2/round3/lab_strength/audit/scripts/extract_canonical.py",
            "track_b/.venv/bin/python track_b/research/q2/round3/lab_strength/audit/scripts/audit_workbooks.py"],
        no_workbook_authoring=True,no_installs=True,no_git_commits=True)
    write_json(RESULTS/"audit_report.json",report)
    files=[dict(path=p.relative_to(ROOT).as_posix(),bytes=p.stat().st_size,sha256=sha(p.read_bytes())) for p in sorted(RESULTS.rglob('*')) if p.is_file() and p.name!='output_index.json']
    write_json(RESULTS/"output_index.json",dict(files=files,self_entry=False))
    print(json.dumps(dict(validation=validation,domain=domain,report_sha256=sha((RESULTS/'audit_report.json').read_bytes())),ensure_ascii=False))


if __name__=="__main__":main()
