#!/usr/bin/env python3
"""Read-only inventory and coverage audit for Track B temporal attachments.

AI assistance: Codex, OpenAI, 2026-10-23; exact model/version not exposed.
Generated audit code and charts require team review; see track_b/AI_USAGE.md.
No imputation, physical-range rejection, QC filtering, modelling, or raw-data writes.
Python 3.10+, pandas, numpy, openpyxl, matplotlib. Optional macOS textutil extracts source .doc.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import re
import subprocess
import sys
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import numpy as np
import openpyxl
import pandas as pd


HOURLY_DIR = "04_B赛道/01_逐小时地气观测（2021-2026）"
DAILY_DIR = "04_B赛道/02_逐日平均气温（2005-2025）"
METADATA_URL = "https://gwfnet.net/Metadata/Index/T-2026-08-28-p1f3e5qIqTEyibfPAp3smBeA"
MISSING_SENTINEL = 9999.9


def sha256(path: Path) -> str:
    h = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def iso(value):
    return None if pd.isna(value) else pd.Timestamp(value).isoformat(sep=" ")


def number(value):
    return None if pd.isna(value) else float(value)


def count_runs(mask: pd.Series, times: pd.Series, minimum=1):
    """Contiguous position runs, reported in file order (time issues separately)."""
    runs = []
    indices = np.flatnonzero(mask.to_numpy(dtype=bool))
    if not len(indices):
        return runs
    groups = np.split(indices, np.flatnonzero(np.diff(indices) != 1) + 1)
    for group in groups:
        if len(group) >= minimum:
            runs.append({"start": iso(times.iloc[group[0]]), "end": iso(times.iloc[group[-1]]), "records": int(len(group))})
    return runs


def time_summary(times: pd.Series, freq: str, expected_year: int | None = None):
    present = pd.DatetimeIndex(times.dropna().unique()).sort_values()
    result = {
        "rows": len(times), "unparseable_times": int(times.isna().sum()),
        "duplicate_time_rows": int(times.duplicated(keep=False).sum()),
        "extra_duplicate_rows": int(times.duplicated().sum()),
        "monotonic_increasing": bool(times.is_monotonic_increasing),
        "time_start": iso(times.min()), "time_end": iso(times.max()),
        "unique_times": len(present),
        "missing_between_min_max": len(pd.date_range(present.min(), present.max(), freq=freq).difference(present)) if len(present) else None,
        "nonstandard_step_rows": int((times.diff().dropna() != pd.Timedelta(hours=1 if freq == "h" else 24)).sum()),
    }
    if expected_year is not None:
        begin = pd.Timestamp(year=expected_year, month=1, day=1)
        end = pd.Timestamp(year=expected_year + 1, month=1, day=1)
        expected = pd.date_range(begin, end, freq=freq, inclusive="left")
        result.update(expected_full_year_records=len(expected), missing_full_year_timestamps=len(expected.difference(present)),
                      outside_filename_year_rows=int(((times < begin) | (times >= end)).sum()))
    return result


def raw_time_format(value):
    text = str(value)
    if re.fullmatch(r"\d{4}-\d{1,2}-\d{1,2} \d{1,2}:\d{2}", text):
        return "Y-M-D H:mm"
    if re.fullmatch(r"\d{4}/\d{1,2}/\d{1,2} \d{1,2}:\d{2}", text):
        return "Y/M/D H:mm"
    return "other"


def clean_measurement(raw):
    numeric = pd.to_numeric(raw, errors="coerce")
    finite = pd.Series(np.isfinite(numeric.to_numpy(dtype=float)), index=raw.index)
    sentinel = numeric.eq(MISSING_SENTINEL)
    # Presence != quality. A finite non-sentinel QC=2 value still counts as present.
    present = numeric.notna() & finite & ~sentinel
    return numeric, sentinel, present


def variable_summary(raw, times, base):
    numeric, sentinel, present = clean_measurement(raw)
    missing_runs = count_runs(~present, times)
    values = numeric[present]
    return {
        **base, "records": len(raw), "raw_blank_or_na": int(raw.isna().sum()),
        "nonnumeric_nonblank": int((raw.notna() & numeric.isna()).sum()),
        "nonfinite_numeric": int((numeric.notna() & ~np.isfinite(numeric)).sum()),
        "sentinel_9999_9": int(sentinel.sum()), "present_non_sentinel": int(present.sum()),
        "missing_or_unusable_numeric": int((~present).sum()),
        "presence_rate": float(present.mean()) if len(raw) else None,
        "first_present": iso(times[present].min()), "last_present": iso(times[present].max()),
        "minimum_present": number(values.min()), "maximum_present": number(values.max()),
        "longest_missing_run_records": max((r["records"] for r in missing_runs), default=0),
        "missing_run_count": len(missing_runs),
    }, present


def write_csv(out: Path, name: str, rows: list[dict], columns=None):
    table = pd.DataFrame(rows, columns=columns if not rows else None)
    table.to_csv(out / name, index=False, encoding="utf-8-sig", lineterminator="\n")


def audit_hourly(root: Path, out: Path, event_date: str):
    files, variables, qc_counts, qc_cross, monthly, event_rows, event_daily, long_runs, extrema = [], [], [], [], [], [], [], [], []
    by_kind = {}
    event = pd.Timestamp(event_date)
    windows = {
        "antecedent_30d": (event - pd.Timedelta(days=30), event),
        "antecedent_7d": (event - pd.Timedelta(days=7), event),
        "event_calendar_day": (event, event + pd.Timedelta(days=1)),
        "following_calendar_day": (event + pd.Timedelta(days=1), event + pd.Timedelta(days=2)),
    }
    for path in sorted((root / HOURLY_DIR).glob("*.csv")):
        kind, _, year_text = path.stem.split("_")
        year = int(year_text)
        data = pd.read_csv(path, encoding="utf-8-sig", low_memory=False)
        raw_times = data["Timestamp"]
        times = pd.to_datetime(raw_times, format="mixed", errors="coerce")
        measurements = [col for col in data.columns if col != "Timestamp" and not col.startswith("QC_")]
        qc_columns = [col for col in data.columns if col.startswith("QC_")]
        all_present = pd.DataFrame({col: clean_measurement(data[col])[2] for col in measurements})
        files.append({
            "file": str(path.relative_to(root)), "kind": kind, "filename_year": year,
            **time_summary(times, "h", year), "columns": len(data.columns),
            "measurement_columns": len(measurements), "qc_columns": len(qc_columns),
            "timestamp_format_counts": json.dumps(raw_times.map(raw_time_format).value_counts().to_dict(), ensure_ascii=False),
            "any_measurement_present_rows": int(all_present.any(axis=1).sum()),
            "all_measurements_present_rows": int(all_present.all(axis=1).sum()),
            "all_measurements_missing_rows": int((~all_present.any(axis=1)).sum()),
            "first_any_measurement_present": iso(times[all_present.any(axis=1)].min()),
            "last_any_measurement_present": iso(times[all_present.any(axis=1)].max()),
            "raw_timezone_suffix_present_rows": int(raw_times.astype(str).str.contains(r"Z$|[+-]\d{2}:\d{2}$", regex=True).sum()),
            "negative_measurement_values_before_qc_filter": int(sum((clean_measurement(data[col])[0][all_present[col]] < 0).sum() for col in measurements)),
        })
        by_kind.setdefault(kind, []).append((data, times))
        for col in qc_columns:
            counts = data[col].value_counts(dropna=False)
            for value, count in counts.items():
                qc_counts.append({"file": path.name, "kind": kind, "year": year, "qc_column": col,
                                  "qc_raw_value": "<NA>" if pd.isna(value) else str(value), "count": int(count)})
        for position, col in enumerate(data.columns):
            if col not in measurements:
                continue
            base = {"file": path.name, "kind": kind, "year": year, "variable": col}
            summary, present = variable_summary(data[col], times, base)
            qc_col = data.columns[position + 1] if position + 1 < len(data.columns) and data.columns[position + 1].startswith("QC_") else None
            summary["qc_column"] = qc_col
            summary["qc_link_rule"] = "immediately_following_QC_column" if qc_col else "no_adjacent_QC_column"
            if qc_col:
                qc_numeric = pd.to_numeric(data[qc_col], errors="coerce")
                for code in [0, 1, 2, 4, 8, 9]:
                    summary[f"present_qc_{code}"] = int((present & qc_numeric.eq(code)).sum())
                cross = pd.crosstab(data[qc_col].astype(str), np.where(present, "present_non_sentinel", "missing_or_unusable_numeric"))
                for code in cross.index:
                    for state in cross.columns:
                        qc_cross.append({**base, "qc_column": qc_col, "qc_raw_value": code, "numeric_state": state, "count": int(cross.loc[code, state])})
            variables.append(summary)
            numeric = pd.to_numeric(data[col], errors="coerce")
            if present.any():
                for label, idx in [("minimum", numeric[present].idxmin()), ("maximum", numeric[present].idxmax())]:
                    extrema.append({**base, "extreme": label, "time": iso(times.loc[idx]), "value": float(numeric.loc[idx]),
                                    "source_csv_row_1based_including_header": int(idx) + 2,
                                    "qc_column": qc_col, "qc_raw_value": str(data.loc[idx, qc_col]) if qc_col else None})
            for run in count_runs(~present, times, minimum=6):
                long_runs.append({**base, **run})
            for month in range(1, 13):
                month_start = pd.Timestamp(year=year, month=month, day=1)
                month_end = month_start + pd.offsets.MonthBegin(1)
                select = (times >= month_start) & (times < month_end)
                expected = len(pd.date_range(month_start, month_end, freq="h", inclusive="left"))
                monthly.append({**base, "month": month, "expected_hours": expected, "rows": int(select.sum()),
                                "present_hours": int((select & present).sum()), "presence_rate_expected_hours": float((select & present).sum() / expected)})
            # Evaluate only files that can intersect the requested event windows.
            if year == event.year:
                for label, (start, end) in windows.items():
                    select = (times >= start) & (times < end)
                    expected = int((end - start).total_seconds() / 3600)
                    row = {**base, "window": label, "window_start_inclusive": iso(start), "window_end_exclusive": iso(end),
                           "expected_hours": expected, "timestamp_rows": int(select.sum()),
                           "present_hours": int((select & present).sum()), "presence_rate_expected_hours": float((select & present).sum() / expected)}
                    if qc_col:
                        for code in [0, 1, 2, 4, 8, 9]:
                            row[f"present_qc_{code}"] = int((select & present & qc_numeric.eq(code)).sum())
                    event_rows.append(row)
                for day in pd.date_range(event - pd.Timedelta(days=30), event + pd.Timedelta(days=1), freq="D"):
                    select = (times >= day) & (times < day + pd.Timedelta(days=1))
                    event_daily.append({**base, "date": str(day.date()), "expected_hours": 24,
                                        "present_hours": int((select & present).sum()),
                                        "qc0_present_hours": int((select & present & qc_numeric.eq(0)).sum()) if qc_col else None})
    combined = {}
    for kind, yearly in by_kind.items():
        data = pd.concat([item[0] for item in yearly], ignore_index=True)
        times = pd.concat([item[1] for item in yearly], ignore_index=True)
        measurements = [col for col in data.columns if col != "Timestamp" and not col.startswith("QC_")]
        presence = pd.DataFrame({col: clean_measurement(data[col])[2] for col in measurements})
        combined[kind] = {**time_summary(times, "h"), "first_any_measurement_present": iso(times[presence.any(axis=1)].min()),
                          "last_any_measurement_present": iso(times[presence.any(axis=1)].max()),
                          "any_measurement_present_rows": int(presence.any(axis=1).sum()),
                          "all_measurements_missing_rows": int((~presence.any(axis=1)).sum()),
                          "all_measurements_present_rows": int(presence.all(axis=1).sum()),
                          "measurement_columns": measurements,
                          "no_adjacent_qc_columns": [col for col in measurements if col == data.columns[-1] or not data.columns[data.columns.get_loc(col) + 1].startswith("QC_")]}
    write_csv(out, "hourly_files.csv", files)
    write_csv(out, "hourly_variables.csv", variables)
    write_csv(out, "hourly_qc_distribution.csv", qc_counts)
    write_csv(out, "hourly_qc_value_crosstab.csv", qc_cross)
    write_csv(out, "hourly_monthly_coverage.csv", monthly)
    write_csv(out, "hourly_event_windows.csv", event_rows)
    write_csv(out, "hourly_event_daily_coverage.csv", event_daily)
    write_csv(out, "hourly_missing_runs_ge6h.csv", long_runs)
    write_csv(out, "hourly_extrema.csv", extrema)
    return combined


def audit_daily(root: Path, out: Path):
    files, variables, monthly, combined_tables = [], [], [], []
    for path in sorted((root / DAILY_DIR).glob("*.xlsx")):
        year = int(re.search(r"(\d{4})\.xlsx$", path.name).group(1))
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
        for sheet in workbook:
            formula_cells = sum(cell.data_type == "f" for row in sheet.iter_rows() for cell in row)
            data = pd.read_excel(path, sheet_name=sheet.title, engine="openpyxl")
            if "时间" not in data.columns:
                raise ValueError(f"Missing 时间 column: {path} / {sheet.title}")
            times = pd.to_datetime(data["时间"], errors="coerce")
            combined_tables.append(data)
            files.append({"file": str(path.relative_to(root)), "sheet": sheet.title, "filename_year": year,
                          **time_summary(times, "D", year), "columns": len(data.columns), "column_names": json.dumps(list(data.columns), ensure_ascii=False),
                          "formula_cells": formula_cells,
                          "non_midnight_rows": int((times.dropna() != times.dropna().dt.normalize()).sum())})
            for col in data.columns:
                if col == "时间":
                    continue
                base = {"file": path.name, "sheet": sheet.title, "year": year, "variable": col}
                summary, present = variable_summary(data[col], times, base)
                summary["sentinel_definition_status"] = "9999.9_counted_for_audit_only_not_documented_for_daily_dataset"
                variables.append(summary)
                for month in range(1, 13):
                    start = pd.Timestamp(year=year, month=month, day=1)
                    end = start + pd.offsets.MonthBegin(1)
                    select = (times >= start) & (times < end)
                    expected = (end - start).days
                    monthly.append({**base, "month": month, "expected_days": expected, "timestamp_rows": int(select.sum()),
                                    "present_days": int((select & present).sum()), "presence_rate_expected_days": float((select & present).sum() / expected)})
        workbook.close()
    all_data = pd.concat(combined_tables, ignore_index=True)
    all_times = pd.to_datetime(all_data["时间"], errors="coerce")
    combined = {**time_summary(all_times, "D"), "counties": list(all_data.columns[1:]), "file_count": len(files),
                "formula_cells_total": sum(item["formula_cells"] for item in files),
                "missing_dates_between_min_max": [str(ts.date()) for ts in pd.date_range(all_times.min(), all_times.max(), freq="D").difference(pd.DatetimeIndex(all_times))],
                "timezone": "not_explicit_in_local_xlsx_or_source_doc", "unit_in_xlsx_headers": "not_explicit", "spatial_aggregation_method": "not_explicit_in_local_xlsx_or_source_doc"}
    write_csv(out, "daily_workbooks.csv", files)
    write_csv(out, "daily_variables.csv", variables)
    write_csv(out, "daily_monthly_coverage.csv", monthly)
    source_doc = root / DAILY_DIR / "数据引用-数据来源.doc"
    try:
        extracted = subprocess.run(["textutil", "-convert", "txt", "-stdout", str(source_doc)], capture_output=True, check=True, text=True).stdout
        (out / "daily_source_doc_extracted.txt").write_text(extracted, encoding="utf-8")
        combined["source_doc_extraction"] = "macOS textutil; see daily_source_doc_extracted.txt"
    except (FileNotFoundError, subprocess.CalledProcessError):
        combined["source_doc_extraction"] = "unavailable: textutil missing or extraction failed"
    return combined


def plot_quality_overview(out: Path, summary):
    """Coverage-only diagnostic. Equal-weight mean across each type's variables."""
    monthly = pd.read_csv(out / "hourly_monthly_coverage.csv")
    monthly["date"] = pd.to_datetime(dict(year=monthly.year, month=monthly.month, day=1))
    months = pd.date_range(monthly.date.min(), monthly.date.max(), freq="MS")
    kinds = ["MET", "RADM", "SOIL", "FLUX"]
    table = monthly.groupby(["kind", "date"])["presence_rate_expected_hours"].mean().unstack("date").reindex(index=kinds, columns=months)
    fig = plt.figure(figsize=(13.8, 5.8), facecolor="#fafaf7")
    ax = fig.add_axes([0.14, 0.37, 0.74, 0.39])
    mesh = ax.imshow(table.to_numpy(), vmin=0, vmax=1, aspect="auto", cmap="cividis", interpolation="nearest")
    ax.set_yticks(range(len(kinds)), labels=[f"{kind} ({len(summary['hourly'][kind]['measurement_columns'])} vars)" for kind in kinds])
    ax.tick_params(axis="y", length=0, labelsize=11)
    year_starts = [i for i, month in enumerate(months) if month.month == 1]
    ax.set_xticks([i + 5.5 for i in year_starts], labels=[str(months[i].year) for i in year_starts], fontsize=11)
    ax.tick_params(axis="x", length=0, pad=8)
    for i in year_starts[1:]:
        ax.axvline(i - 0.5, color="#fafaf7", linewidth=1.3)
    for y in [0.5, 1.5, 2.5]:
        ax.axhline(y, color="#fafaf7", linewidth=1.2)
    # Number months 01/04/07/10 only, to keep the 72-cell audit readable.
    minor = [i for i, month in enumerate(months) if month.month in [1, 4, 7, 10]]
    top = ax.secondary_xaxis("top")
    top.set_xticks(minor, labels=[f"{months[i].month:02}" for i in minor], fontsize=8)
    top.tick_params(length=0, pad=4)
    for spine in list(ax.spines.values()) + list(top.spines.values()):
        spine.set_visible(False)
    event_month = pd.Timestamp(summary["event_date_scenario"]).to_period("M").to_timestamp()
    event_index = months.get_indexer([event_month])[0]
    if event_index >= 0:
        ax.add_patch(Rectangle((event_index - 0.5, -0.5), 1, len(kinds), fill=False, edgecolor="#d54433", linewidth=1.8))
    cax = fig.add_axes([0.90, 0.37, 0.012, 0.39])
    colorbar = fig.colorbar(mesh, cax=cax, ticks=[0, 0.25, 0.5, 0.75, 1])
    colorbar.ax.set_yticklabels(["0%", "25%", "50%", "75%", "100%"])
    colorbar.outline.set_visible(False)
    colorbar.ax.tick_params(length=0, labelsize=9)
    fig.text(0.08, 0.92, "Temporal data coverage", fontsize=23, weight="bold", color="#19252a")
    fig.text(0.08, 0.855, "Monthly mean of variable presence rates | finite values excluding 9999.9 | denominator: full calendar-month hours", fontsize=10.5, color="#4c575b")
    fig.text(0.08, 0.265, "Presence is not QC acceptance. QC=1/2/4/9 values remain counted; FLUX and precipitation have no QC columns.", fontsize=10.2, color="#28383e")
    fig.text(0.08, 0.205, "Red outline: Aug 2026 scenario month. Last non-sentinel record: 2026-08-27 08:00 (source clock).", fontsize=10.2, color="#28383e")
    fig.text(0.08, 0.145, "Sep-Dec 2026 contain timestamp rows but no non-sentinel measurements. The 2021 leading grid also contains placeholders.", fontsize=10.2, color="#28383e")
    fig.text(0.08, 0.065, "Daily ERA5-Land-derived tables: 7,519 continuous dates, 5 counties; 2005-01-01 to 2025-08-02. No 2026 daily table.", fontsize=10.2, color="#28383e")
    fig.savefig(out / "quality_overview.png", dpi=180, facecolor=fig.get_facecolor())
    plt.close(fig)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--root", type=Path, default=Path(__file__).resolve().parents[2])
    parser.add_argument("--out", type=Path, default=None)
    parser.add_argument("--event-date", default="2026-08-26", help="scenario calendar date, not an asserted observed onset time")
    parser.add_argument("--audit-date", default="2026-10-23", help="user-local date of source access / audit")
    args = parser.parse_args()
    root = args.root.resolve()
    out = (args.out or root / "track_b/audit/temporal").resolve()
    raw_dirs = [(root / HOURLY_DIR).resolve(), (root / DAILY_DIR).resolve()]
    if any(out == p or p in out.parents for p in raw_dirs):
        raise ValueError("Output must not be inside a raw-data directory")
    out.mkdir(parents=True, exist_ok=True)
    manifest = [{"file": str(path.relative_to(root)), "bytes": path.stat().st_size, "sha256": sha256(path), "suffix": path.suffix.lower()}
                for directory in raw_dirs for path in sorted(directory.iterdir()) if path.is_file() and not path.name.startswith(".")]
    write_csv(out, "manifest.csv", manifest)
    summary = {
        "audit_date_user_local": args.audit_date, "event_date_scenario": args.event_date,
        "runtime": {"python": platform.python_version(), "executable": sys.executable, "pandas": pd.__version__, "numpy": np.__version__, "openpyxl": openpyxl.__version__, "matplotlib": matplotlib.__version__},
        "script_sha256": sha256(Path(__file__)),
        "definitions": {
            "present_non_sentinel": "finite numeric value unequal to 9999.9; not proof of QC acceptance or physical validity",
            "hourly_missing_sentinel": MISSING_SENTINEL,
            "hourly_sentinel_source": METADATA_URL,
            "qc_link_rule": "immediately following QC_ column, preserve actual spelling; no invented QC for Precip or FLUX",
            "qc_policy": "raw distributions and numeric-presence cross-tabs only; no records removed by QC",
            "time_policy": "parse naive timestamps, report metadata UTC+8 separately; no UTC conversion, resampling or inferred event onset",
            "event_windows": "half-open calendar windows; day starts at source-local 00:00 if hourly metadata applies",
            "missing_run": "consecutive file-position records absent by numeric-presence definition; runs >=6 records exported",
            "daily_sentinel": "9999.9 counted as audit marker candidate; no daily source confirmation; count is zero in this run if reported so",
        },
        "metadata_assertions": {
            "source_url": METADATA_URL, "access_date_user_local": args.audit_date,
            "hourly_timezone": "UTC+8", "claimed_start": "2021-04-26", "claimed_end": "2026-08-27",
            "qc_codes": {"0": "correct", "1": "suspicious", "2": "error", "4": "interpolated", "8": "missing", "9": "not_quality_controlled"},
            "station_latitude": 28.86, "station_longitude": 85.29, "station_elevation_m": 4140,
            "stated_distance_to_affected_port_km": 65,
            "metadata_version_local_equivalence": "not verified: mirror filenames differ from local names; no remote file hash comparison",
        },
        "manifest_file_count": len(manifest), "hourly": audit_hourly(root, out, args.event_date), "daily": audit_daily(root, out),
        "raw_hashes_unchanged_after_audit": all(sha256(root / item["file"]) == item["sha256"] for item in manifest),
    }
    with (out / "summary.json").open("w", encoding="utf-8") as stream:
        json.dump(summary, stream, ensure_ascii=False, indent=2, allow_nan=False)
        stream.write("\n")
    plot_quality_overview(out, summary)
    outputs = [{"file": p.name, "bytes": p.stat().st_size, "sha256": sha256(p)} for p in sorted(out.iterdir()) if p.is_file() and p.name != "output_manifest.json"]
    (out / "output_manifest.json").write_text(json.dumps(outputs, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"output": str(out), "raw_files": len(manifest), "hourly_types": list(summary["hourly"]), "daily_time_start": summary["daily"]["time_start"], "daily_time_end": summary["daily"]["time_end"], "raw_hashes_unchanged": summary["raw_hashes_unchanged_after_audit"]}, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
