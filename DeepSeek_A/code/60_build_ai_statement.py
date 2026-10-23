# -*- coding: utf-8 -*-
"""60_build_ai_statement.py —— 生成《AI工具使用详情说明》DOCX（按竞赛 AI 规定要求）"""
import os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

BASE = r"C:\Users\17601\Desktop\MathorCup\DeepSeek_A"
PAPER = os.path.join(BASE, 'paper')
os.makedirs(PAPER, exist_ok=True)
OUT = os.path.join(PAPER, 'AI工具使用详情说明.docx')

doc = Document()
st = doc.styles['Normal']
st.font.name = 'Times New Roman'
st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
for sec in doc.sections:
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.left_margin = Cm(2.5); sec.right_margin = Cm(2.5)
    sec.top_margin = Cm(2.3); sec.bottom_margin = Cm(2.2)



def srun(run, size=10.5, bold=False, cn='宋体'):
    run.font.name = 'Times New Roman'
    run.font.size = Pt(size)
    run.bold = bold
    run._element.get_or_add_rPr()
    run._element.rPr.get_or_add_rFonts().set(qn('w:eastAsia'), cn)
    return run

def h(text, size=14, before=10, after=5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before); p.paragraph_format.space_after = Pt(after)
    r = srun(p.add_run(text), size=size, bold=True, cn='黑体')
    return p


def par(text, size=10.5, indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    r = srun(p.add_run(text), size=size, cn='宋体')
    return p


p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = srun(p.add_run('AI 工具使用详情说明'), size=16, bold=True, cn='黑体')
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = srun(p.add_run('2026 年第七届 MathorCup 数学应用挑战赛——大数据竞赛  赛道 A'), size=10.5)

h('一、所用 AI 工具的名称与版本')
par('工具名称：DeepSeek；版本/型号：DeepSeek-V3.2（深度求索）；使用环境：DeepSeek Harness 编码智能体；'
    '开发机构：杭州深度求索人工智能基础技术研究有限公司（DeepSeek）；使用日期：2026 年 10 月 23 日。')

h('二、具体使用目的与应用环节')
par('本参赛队在竞赛过程中将 DeepSeek 作为辅助工具使用，具体环节包括：')
par('1. 数据处理与分析脚本的辅助编写与调试：包括数据文件读取与校验、缺失值处理、秩-高斯变换、'
    '交叉验证循环、bootstrap 重采样、绘图脚本等。所有脚本由参赛队逐行运行验证，并对其中错误进行了人工修正。')
par('2. 建模方案的辅助设计：包括候选模型清单、评价指标与稳定性检验方案、缺陷约束投影的公式推导等。'
    '最终采用的方法、超参数与验证方案由参赛队讨论确定。')
par('3. 图表绘制：论文中的 15 幅图由脚本自动生成，DeepSeek 辅助编写了绘图代码；'
    '图的设计与数据口径由参赛队确定。')
par('4. 论文文字润色与排版：DeepSeek 辅助生成论文初稿的部分段落文字与 DOCX 排版脚本，'
    '参赛队对全部文字进行了事实核对、数据核对与改写，确保结论与计算结果一致。')

h('三、关键交互记录（摘要）')
par('（1）提示词：请读取附件 1—8，核验 1.csv 的数值是否为对数变换、底数是 log10 还是 ln，并给出证据。'
    '回复要点：通过 8.xlsx 中 Calibrated chemical dataset 的浓度与 1.csv 数值建立回归，'
    '斜率 ln(10)=2.303 可判定底数为 10；37 个参数列（含 pH、乙醇、色度、苦度、CO2）均为 log10 值。'
    '采纳情况：采纳，并在论文 2.2 节报告了核验结果与示例表。')
par('（2）提示词：4.csv 的感官数值有负值且不是 0—6 分，如何给出题目要求的 0—6 分级评分？'
    '回复要点：推断其为逐品评员 z 化后的均值，可建立 y≈a+bX 的线性潜变量模型，'
    '用等距高斯混合 + EM 反演等级，并用仿真评估可恢复精度。采纳情况：采纳其建模框架，'
    '参赛队补充了两阶段网格搜索实现、仿真实验设计（不同品评员人数/噪声）与灵敏度分析。')
par('（3）提示词：每个感官维度的 Top-3 关键化合物如何兼顾稳定性与化学合理性？'
    '回复要点：建议采用"单变量初筛 + 多方法秩聚合 + bootstrap 入选频率 + 方法间 Jaccard"方案，'
    '并用标志化合物命中率与来源一致性检验合理性。采纳情况：采纳，并据文献补充了标志化合物先验清单。')
par('（4）提示词：如何保证跨风格迁移时缺陷化合物不恶化？回复要点：在输出空间对缺陷化合物浓度做'
    '"上限投影"（上限取源样本与目标风格 75% 分位的较小者），可精确保证缺陷指数不高于源样本。'
    '采纳情况：采纳，并由参赛队推导了 log10/原始浓度两种情形下的精确投影公式。')
par('（5）提示词：请检查论文中"pH/乙醇保留原始量纲"这一表述是否正确。'
    '回复要点：与 1.csv 实际数据不符，1.csv 中这些列同样为 log10 变换值。'
    '采纳情况：采纳，论文中已澄清该点并在数据核验一节给出证据。')

h('四、采纳内容与人工修改情况')
par('1. 代码：DeepSeek 辅助编写的代码全部经参赛队实际运行、检查中间结果并修正，'
    '代码注释中已标注 AI 辅助情况（工具名称、版本、开发机构、使用日期）。')
par('2. 结果：所有数值结果均由参赛队亲自运行脚本获得，论文中报告的数字与支撑材料中的结果文件'
    '（CSV/JSON）一一对应，未直接采用 AI 生成但未经运行验证的数值。')
par('3. 文字：AI 辅助生成的文字均经参赛队改写与核对，确保符合参赛队对方法与结果的理解；'
    '所有引用文献均经参赛队核查来源。')
par('4. 论文中 AI 辅助内容已在参考文献与文末"AI 工具使用声明"中标注。')

h('五、声明')
par('本参赛队对使用 AI 工具生成内容的正确性、可靠性与学术规范性负责。'
    '本说明与论文正文中的声明一致，并随支撑材料一并提交。')

doc.save(OUT)
print('saved', OUT)
