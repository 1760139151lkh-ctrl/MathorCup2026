# -*- coding: utf-8 -*-
"""论文排版器：把 paper_content_*.py 中的内容块渲染为符合 MathorCup 格式规范的 DOCX。

格式要点：
- 第一页：题目 + 摘要 + 关键词；第二页：目录；之后正文从第 1 页开始编页（页脚居中）；
- 无页眉；全文中文；正文 ≤30 页，附录不限；
- 图表自动编号（内容块中给出编号与标题），图片居中、图题在下，表题在上；
- 目录页码由 toc_pages.json 提供（第一次生成时为占位，后续 update_toc 回填）。
"""
import json
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER, WD_LINE_SPACING, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
import labels as L
import paper_content_1 as C1
import paper_content_2 as C2
import paper_content_3 as C3

sys.stdout.reconfigure(encoding="utf-8")
WORK = L.WORK
PAPER = L.PAPER
FIGS = L.FIGURES
TOC_JSON = PAPER / "toc_pages.json"
EQ_MANIFEST = FIGS / "_equation_manifest.json"
EQ = json.loads(EQ_MANIFEST.read_text(encoding="utf-8")) if EQ_MANIFEST.exists() else {}

CN_BODY = "宋体"
CN_HEAD = "黑体"
EN_BODY = "Times New Roman"
EN_MATH = "Cambria Math"
EN_CODE = "Consolas"

BLOCKS = C1.BLOCKS + C2.BLOCKS + C3.BLOCKS
META = C1.META


# ------------------------------------------------------------------ 基础工具
def set_run_font(run, cn=CN_BODY, en=EN_BODY, size=12, bold=False, color=(0, 0, 0)):
    run.font.name = en
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(*color)
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)


def style_font(style, cn=CN_BODY, en=EN_BODY, size=12, bold=False, color=(0, 0, 0)):
    style.font.name = en
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(*color)
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)


def add_field(paragraph, instr):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar"); fld_begin.set(qn("w:fldCharType"), "begin")
    instr_el = OxmlElement("w:instrText"); instr_el.set(qn("xml:space"), "preserve"); instr_el.text = instr
    fld_sep = OxmlElement("w:fldChar"); fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t"); txt.text = "1"
    fld_end = OxmlElement("w:fldChar"); fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin); run._r.append(instr_el); run._r.append(fld_sep)
    run._r.append(txt); run._r.append(fld_end)


def shade_cell(cell, hex_color="DCE6F1"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def set_cell_margins(table, top=40, bottom=40, left=70, right=70):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{k}"); el.set(qn("w:w"), str(v)); el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def set_table_layout(table, widths):
    """强制固定列宽（LibreOffice/Word 都按此渲染）。"""
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = OxmlElement("w:tblLayout"); layout.set(qn("w:type"), "fixed"); tblPr.append(layout)
    total = sum(int(w.twips) for w in widths)
    tblW = OxmlElement("w:tblW"); tblW.set(qn("w:w"), str(total)); tblW.set(qn("w:type"), "dxa")
    tblPr.append(tblW)
    for j, w in enumerate(widths):
        table.columns[j].width = w
        for row in table.rows:
            row.cells[j].width = w


def set_rows_cant_split(table):
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true")
        trPr.append(el)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true")
    trPr.append(el)


# ------------------------------------------------------------------ 文档骨架
def setup_document():
    doc = Document()
    for sec in doc.sections:
        sec.page_width = Cm(21.0)
        sec.page_height = Cm(29.7)
        sec.top_margin = Cm(2.2)
        sec.bottom_margin = Cm(2.2)
        sec.left_margin = Cm(2.2)
        sec.right_margin = Cm(2.2)
    normal = doc.styles["Normal"]
    style_font(normal, cn=CN_BODY, en=EN_BODY, size=12)
    pf = normal.paragraph_format
    pf.line_spacing = 1.28
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    for name, size, before, after in [("Heading 1", 16, 10, 5), ("Heading 2", 13.5, 8, 4),
                                      ("Heading 3", 12, 6, 3)]:
        st = doc.styles[name]
        style_font(st, cn=CN_HEAD, en=EN_BODY, size=size, bold=True)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.3
        st.paragraph_format.keep_with_next = True
    return doc


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def build_front_matter(doc):
    def centered(text, cn=CN_HEAD, size=14, bold=False, space_before=0, space_after=6, en=EN_BODY):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.3
        r = p.add_run(text)
        set_run_font(r, cn=cn, en=en, size=size, bold=bold)
        return p

    centered(META["competition"], cn=CN_HEAD, size=15, bold=True, space_before=6, space_after=4)
    centered(META["track"], cn=CN_HEAD, size=12.5, bold=False, space_after=16)

    # 队伍编号 / 赛道
    table = doc.add_table(rows=2, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = [Cm(4.2), Cm(7.2)]
    vals = [("队伍编号", META["team_id"]), ("赛道", META["track_code"])]
    for i, (a, b) in enumerate(vals):
        for j, text in enumerate((a, b)):
            cell = table.cell(i, j)
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            set_run_font(r, cn=CN_HEAD if j == 0 else CN_BODY, size=11)
    # 边框
    tbl = table._tbl
    tblPr = tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single"); el.set(qn("w:sz"), "4"); el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000"); borders.append(el)
    tblPr.append(borders)

    for line in META["title_lines"]:
        centered(line, cn=CN_HEAD, size=16.5, bold=True, space_before=8, space_after=4)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(4)
    r = p.add_run("摘   要"); set_run_font(r, cn=CN_HEAD, size=15, bold=True)

    for para in META["abstract"]:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Pt(22)
        p.paragraph_format.line_spacing = 1.24
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(para); set_run_font(r, size=11)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run("关键词："); set_run_font(r, cn=CN_HEAD, size=11, bold=True)
    r2 = p.add_run(META["keywords"]); set_run_font(r2, size=11)


def build_toc(doc, entries):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.page_break_before = True
    p.paragraph_format.space_before = Pt(6); p.paragraph_format.space_after = Pt(12)
    r = p.add_run("目    录"); set_run_font(r, cn=CN_HEAD, size=16, bold=True)
    pages = {}
    if TOC_JSON.exists():
        pages = json.loads(TOC_JSON.read_text(encoding="utf-8"))
    right = Cm(15.8)
    for level, title in entries:
        para = doc.add_paragraph()
        pf = para.paragraph_format
        pf.line_spacing = 1.02
        pf.space_after = Pt(0)
        pf.left_indent = Cm(0.0 if level == 1 else 0.8)
        pf.tab_stops.add_tab_stop(right, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        r = para.add_run(title)
        set_run_font(r, cn=CN_HEAD if level == 1 else CN_BODY, size=11 if level == 1 else 9.5,
                     bold=(level == 1))
        r2 = para.add_run("\t" + str(pages.get(title, "0")))
        set_run_font(r2, size=11 if level == 1 else 9.5)


def start_body_section(doc):
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.2)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    # 页码从 1 开始
    sectPr = sec._sectPr
    pg = OxmlElement("w:pgNumType"); pg.set(qn("w:start"), "1"); sectPr.append(pg)
    # 页脚居中 PAGE 域
    sec.footer.is_linked_to_previous = False
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(fp, " PAGE ")
    for r in fp.runs:
        set_run_font(r, size=10.5)
    return sec


# ------------------------------------------------------------------ 正文块
def add_heading(doc, text, level):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    h = doc.add_paragraph(style=style)
    r = h.add_run(text)
    set_run_font(r, cn=CN_HEAD, size={1: 16, 2: 13.5, 3: 12}[level], bold=True)
    if level == 1:
        h.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return h


def add_paragraph(doc, text, indent=True, size=12, align="justify", space_after=0, bold=False):
    p = doc.add_paragraph()
    p.alignment = {"justify": WD_ALIGN_PARAGRAPH.JUSTIFY, "center": WD_ALIGN_PARAGRAPH.CENTER,
                   "left": WD_ALIGN_PARAGRAPH.LEFT}[align]
    if indent:
        p.paragraph_format.first_line_indent = Pt(size * 2)
    p.paragraph_format.line_spacing = 1.28
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    set_run_font(r, size=size, bold=bold)
    return p


def add_list(doc, items, size=12):
    for it in items:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.74)
        p.paragraph_format.first_line_indent = Cm(-0.0)
        p.paragraph_format.line_spacing = 1.28
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(it)
        set_run_font(r, size=size)


def add_equation(doc, text, number=""):
    info = EQ.get(str(number))
    if info is None:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text + (f"    ({number})" if number else ""))
        set_run_font(r, cn=CN_BODY, en=EN_MATH, size=11.5)
        return
    tab = doc.add_table(rows=1, cols=2)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    w = min(info["width_cm"], 14.6)
    c0, c1 = tab.cell(0, 0), tab.cell(0, 1)
    set_table_layout(tab, [Cm(w + 0.4), Cm(1.7)])
    p0 = c0.paragraphs[0]
    p0.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p0.paragraph_format.space_before = Pt(3); p0.paragraph_format.space_after = Pt(3)
    p0.add_run().add_picture(str(FIGS / "equations" / info["file"]), width=Cm(w))
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.paragraph_format.space_before = Pt(3); p1.paragraph_format.space_after = Pt(3)
    r = p1.add_run(f"({number})"); set_run_font(r, size=10.5)


def add_figure(doc, filename, caption, width_cm=15.2):
    path = FIGS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    run.add_picture(str(path), width=Cm(width_cm * 0.86))
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    cap.paragraph_format.line_spacing = 1.15
    r = cap.add_run(caption)
    set_run_font(r, cn=CN_HEAD, size=10, bold=False)


def add_table(doc, caption, headers, rows, widths=None, font_size=9, align="center"):
    cap = doc.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_before = Pt(8)
    cap.paragraph_format.space_after = Pt(3)
    cap.paragraph_format.keep_with_next = True
    cap.paragraph_format.line_spacing = 1.15
    r = cap.add_run(caption)
    set_run_font(r, cn=CN_HEAD, size=10, bold=False)
    font_size = max(7.2, font_size * 0.90)
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths is None:
        widths = [Cm(15.2 / len(headers))] * len(headers)
    set_cell_margins(table)
    set_table_layout(table, widths)
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.width = widths[j]
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(str(h)); set_run_font(r, cn=CN_HEAD, size=font_size, bold=True)
        shade_cell(cell)
    set_repeat_header(table.rows[0])
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.width = widths[j]
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if align == "center" else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(val)); set_run_font(r, size=font_size)
    set_rows_cant_split(table)
    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)
    sp.paragraph_format.line_spacing = 1.0
    r = sp.add_run(""); set_run_font(r, size=4)


def add_note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    set_run_font(r, size=10)


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph()
        pf = p.paragraph_format
        pf.line_spacing = 1.0
        pf.space_after = Pt(0)
        pf.space_before = Pt(0)
        pf.left_indent = Cm(0.3)
        r = p.add_run(line if line else " ")
        set_run_font(r, cn=CN_BODY, en=EN_CODE, size=8)


def add_reference(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf = p.paragraph_format
    pf.left_indent = Cm(0.85)
    pf.first_line_indent = Cm(-0.85)
    pf.line_spacing = 1.3
    pf.space_after = Pt(2)
    r = p.add_run(text)
    set_run_font(r, size=10.5)


# ------------------------------------------------------------------ 主流程
def collect_toc_entries():
    entries = []
    for b in BLOCKS:
        if b[0] == "h1":
            entries.append((1, b[1]))
        elif b[0] == "h2":
            entries.append((2, b[1]))
    return entries


def build():
    doc = setup_document()
    build_front_matter(doc)
    entries = collect_toc_entries()
    build_toc(doc, entries)
    start_body_section(doc)

    for b in BLOCKS:
        kind = b[0]
        if kind == "h1":
            add_heading(doc, b[1], 1)
        elif kind == "h2":
            add_heading(doc, b[1], 2)
        elif kind == "h3":
            add_heading(doc, b[1], 3)
        elif kind == "p":
            add_paragraph(doc, b[1])
        elif kind == "list":
            add_list(doc, b[1])
        elif kind == "eq":
            add_equation(doc, b[1], b[2] if len(b) > 2 else "")
        elif kind == "fig":
            add_figure(doc, b[1], b[2], b[3] if len(b) > 3 else 15.2)
        elif kind == "table":
            add_table(doc, b[1], b[2], b[3], b[4] if len(b) > 4 else None,
                      b[5] if len(b) > 5 else 9)
        elif kind == "note":
            add_note(doc, b[1])
        elif kind == "code":
            add_code(doc, b[1])
        elif kind == "ref":
            add_reference(doc, b[1])
        elif kind == "pagebreak":
            add_page_break(doc)
        elif kind == "center":
            add_paragraph(doc, b[1], indent=False, size=b[2] if len(b) > 2 else 12,
                          align="center", bold=(b[3] if len(b) > 3 else False))
        else:
            raise ValueError(f"unknown block {kind}")

    out = PAPER / "A赛道论文_啤酒风味复原版.docx"
    doc.save(out)
    print("saved:", out)
    print("TOC entries:", len(entries))
    return out


if __name__ == "__main__":
    build()
