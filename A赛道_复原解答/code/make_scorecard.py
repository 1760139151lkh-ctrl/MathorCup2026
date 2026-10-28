# -*- coding: utf-8 -*-
"""填写官方 3.xlsx 评分表：把问题 1 的 0–6 工作等级（协议 A）写入 12 张评分卡与 summary。

- 保留原模板的版式、合并单元格、未涉及项目的公式；
- 只把本方法给出的 50 个目标分数写到对应分数格（Q7/AG7 等），A-16/F-16（乙醛，无目标列）写 N/A；
- 在每张卡 B104 与 summary A15 写入评分口径说明（不遮挡原模板内容）。
输出：results/Q1_评分表_官方3xlsx格式.xlsx
"""
import sys
from pathlib import Path

import openpyxl
import pandas as pd

sys.path.insert(0, str(Path(__file__).resolve().parent))
import labels as L

sys.stdout.reconfigure(encoding="utf-8")
OUT = L.RESULTS
GRADES = OUT / "Q1_0-6分级_协议A_全拟合.csv"
SRC = L.ORIG / "3.xlsx"
DST = OUT / "Q1_评分表_官方3xlsx格式.xlsx"

grades = pd.read_csv(GRADES)
assert len(grades) == 250
g_index = grades.set_index("beer")

tpl = L.TEMPLATE_MAP
note_card = ("PROTOCOL A（工作评分）: grade = clip(floor(3 + (z-μ)/σ + 0.5), 0, 6); "
             "μ/σ 来自全部 250 个样本的 D04 连续评分；原始 0–6 等级不可唯一识别。")
note_summary = ("本表为问题 1 的 0–6 工作等级结果（协议 A，全拟合模型）。"
                "原始评分经逐品评员标准化后仅能观测到连续值，绝对等级不可唯一还原；"
                "协议 A 为题目要求的可复核、单调分级口径，协议 B 的对照见结果 CSV。")

wb = openpyxl.load_workbook(SRC)
card_sheets = [s for s in wb.sheetnames if s.startswith("beer_")]
assert len(card_sheets) == 12, card_sheets

n_filled = 0
for si, sheet in enumerate(card_sheets, start=1):
    ws = wb[sheet]
    # 每张卡对应 Beer# = 1..12（原题说明：10 个 sheet 对应 beer_1~beer_10；实际 12 张）
    if si not in g_index.index:
        continue
    row = g_index.loc[si]
    for _, item in tpl.iterrows():
        cell = item["score_cell"]
        target = item["target_id"]
        if pd.isna(target) or str(target).strip() == "":
            ws[cell] = "N/A"
            continue
        yid = str(target).strip()
        col = L.Y2NAME[yid]
        ws[cell] = int(row[col])
        n_filled += 1
    ws["B104"] = note_card

# summary：第 2–13 行对应 beer_1..beer_12；列 B–AY 对应 50 个目标
ws = wb["summary"]
header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
col_of = {str(name).strip(): idx + 1 for idx, name in enumerate(header) if name is not None}
# 原模板 summary 表头存在三处拼写差异（4vg / A_ester_fruity / F_ester_fruity），做别名映射
alias = {"4vg": "X4vg", "A_ester_fruity": "A_esters_fruity", "F_ester_fruity": "F_esters_fruity"}
col_of = {alias.get(k, k): v for k, v in col_of.items()}
n_sum = 0
for si in range(1, 13):
    r = si + 1
    beer_cell = ws.cell(row=r, column=1).value
    try:
        beer = int(float(str(beer_cell)))
    except Exception:
        beer = si
    if beer not in g_index.index:
        continue
    row = g_index.loc[beer]
    for yid, name in L.Y2NAME.items():
        c = col_of.get(name)
        if c is None:
            continue
        ws.cell(row=r, column=c).value = int(row[name])
        n_sum += 1
ws.cell(row=15, column=1).value = note_summary

wb.save(DST)
print("saved", DST)
print("卡片分数格写入:", n_filled, " summary 写入:", n_sum)

# 回读核对：卡片分数格与成绩矩阵一致
wb2 = openpyxl.load_workbook(DST, data_only=True)
bad = 0
for si, sheet in enumerate(card_sheets, start=1):
    row = g_index.loc[si]
    ws2 = wb2[sheet]
    for _, item in tpl.iterrows():
        target = item["target_id"]
        if pd.isna(target) or str(target).strip() == "":
            continue
        cell = item["score_cell"]
        if int(ws2[cell].value) != int(row[L.Y2NAME[str(target).strip()]]):
            bad += 1
print("回读不一致格数:", bad)
