"""只读输入的数据侦察；报告证据，不把未识别的标度猜成事实。"""
from pathlib import Path
import json
import re
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'A题_数学手记' / 'results'
OUT.mkdir(parents=True, exist_ok=True)


def main():
    frames = {name: pd.read_csv(ROOT / 'data' / f'{name}.csv')
              for name in ['D01', 'D02', 'D04', 'D05', 'D06', 'D07']}
    x = frames['D01'].set_index('object')
    y = frames['D04'].set_index('object')
    features = [f'V{i:03d}' for i in range(1, 232)]
    targets = [f'Y{i:02d}' for i in range(1, 51)]
    xnum, ynum = x[features], y[targets]
    book = pd.ExcelFile(ROOT / 'data' / 'D08.xlsx')
    sheets = {s: pd.read_excel(book, s) for s in book.sheet_names}
    scale = sheets['ref_scale'].set_index('object')
    summary = {
        'shape': {k: list(v.shape) for k, v in frames.items()},
        'D08_shape': {k: list(v.shape) for k, v in sheets.items()},
        'D01_missing_cells': int(xnum.isna().sum().sum()),
        'D01_missing_rows': int(xnum.isna().any(axis=1).sum()),
        'D01_missing_columns': int(xnum.isna().any().sum()),
        'D04_missing_cells': int(ynum.isna().sum().sum()),
        'D07_missing_cells': int(frames['D07'].iloc[:, 3:].isna().sum().sum()),
        'regime_counts': x['regime'].value_counts().sort_index().to_dict(),
        'D01_D04_object_alignment': bool(x.index.equals(y.index)),
        'D01_D04_object_id_equal_count': int((x['object_id']==y['object_id']).sum()),
        'D01_D04_regime_equal_count': int((x['regime']==y['regime']).sum()),
    }
    desc = xnum.describe().T
    desc['missing'] = xnum.isna().sum()
    desc['negative'] = (xnum < 0).sum()
    desc['zero'] = (xnum == 0).sum()
    # 中位数绝对偏差只负责标记，不自动删掉罕见但真实的观测。
    med = xnum.median()
    mad = (xnum - med).abs().median()
    desc['robust_z_gt_10'] = (((xnum-med).abs() / (1.4826*mad.replace(0, np.nan))) > 10).sum()
    desc.to_csv(OUT / 'data_audit_feature_summary.csv', encoding='utf-8-sig')
    ydesc = ynum.describe().T
    ydesc['unique_values'] = ynum.nunique()
    ydesc['mean_abs_deviation_from_integer'] = (ynum-ynum.round()).abs().mean()
    ydesc.to_csv(OUT / 'data_audit_target_summary.csv', encoding='utf-8-sig')
    rows, discrepancy_rows = [], []
    # 用参考原生值的对数比较：避免原生量跨度大时让一两个巨数支配误差。
    for var in [c for c in scale if c in features]:
        valid = xnum[var].notna() & scale[var].notna()
        a, b = xnum.loc[valid, var].to_numpy(), scale.loc[valid, var].to_numpy()
        denom = np.maximum(np.abs(b), 1e-12)
        losses = {'identity': np.median(np.abs(a-b)/denom)}
        positive = b > 0
        if positive.any():
            losses['log10'] = np.median(np.abs(a[positive] - np.log10(b[positive])))
            losses['ln'] = np.median(np.abs(a[positive] - np.log(b[positive])))
        best = min(losses, key=losses.get)
        row = {'var_id': var, 'n_pairs': int(valid.sum()), 'best_transform': best,
               'best_error': float(losses[best]), 'identity_max_abs_error':float(np.max(np.abs(a-b))),
               **{f'error_{k}':float(v) for k,v in losses.items()}}
        if best == 'log10':
            nonpositive = b <= 0
            floor_native = np.min(b[positive])/2
            row['native_nonpositive_count'] = int(nonpositive.sum())
            row['native_negative_count'] = int((b < 0).sum())
            row['native_min_positive'] = float(np.min(b[positive]))
            row['replacement_floor_native'] = float(floor_native)
            row['nonpositive_floor_max_log_error'] = float(np.max(np.abs(a[nonpositive]-np.log10(floor_native)))) if nonpositive.any() else np.nan
            row['positive_log10_max_error'] = float(np.max(np.abs(a[positive]-np.log10(b[positive]))))
        else:
            paired = pd.DataFrame({'object':xnum.index[valid], 'var_id':var, 'D01_value':a, 'reference_native':b, 'difference':a-b})
            discrep = paired[paired['difference'].abs()>1e-7]
            row['identity_discrepancies_gt_1e_7'] = int(len(discrep))
            discrepancy_rows.extend(discrep.to_dict('records'))
        rows.append(row)
    transforms = pd.DataFrame(rows)
    transforms.to_csv(OUT / 'data_audit_scale_candidates.csv', index=False, encoding='utf-8-sig')
    pd.DataFrame(discrepancy_rows).to_csv(OUT / 'data_audit_scale_discrepancies.csv', index=False, encoding='utf-8-sig')
    summary['scale_best_counts'] = transforms['best_transform'].value_counts().to_dict()
    summary['scale_strict_identified_counts'] = transforms.loc[transforms['best_error'] < 1e-8, 'best_transform'].value_counts().to_dict()
    summary['identity_reference_discrepancies_gt_1e_7'] = len(discrepancy_rows)
    corr_rows = []
    for name, left, right in [('D02', xnum, xnum), ('D05', ynum, ynum), ('D06', xnum, ynum)]:
        given = pd.read_csv(ROOT / 'data' / f'{name}.csv', index_col=0)
        for method in ['pearson','spearman']:
            # D06非方阵，需在同一样本对齐后计算交叉块。
            if name == 'D06':
                together = pd.concat([left, right], axis=1).corr(method=method)
                computed = together.loc[given.index, given.columns]
            else:
                computed = left.corr(method=method).loc[given.index, given.columns]
            dif = (given-computed).to_numpy()
            if name != 'D06':
                dif = dif[np.triu_indices_from(dif, k=1)]
            corr_rows.append({'matrix':name,'method':method,'n_compared':int(np.isfinite(dif).sum()),
                              'MAE':float(np.nanmean(np.abs(dif))),'max_abs_error':float(np.nanmax(np.abs(dif))),
                              'RMSE':float(np.sqrt(np.nanmean(dif**2)))})
    pd.DataFrame(corr_rows).to_csv(OUT / 'data_audit_correlation_check.csv', index=False, encoding='utf-8-sig')
    # 公式与缓存分开读；空白模板中的公式0不是评分者实际给的0。
    wb = load_workbook(ROOT / 'data' / 'D03.xlsx', data_only=False)
    cached = load_workbook(ROOT / 'data' / 'D03.xlsx', data_only=True)
    card_rows = []
    scoring_rows = {'Q':[7,10,12,14,16,20,23,25,27,29,33,36,38,40,42,45,50,54,58,62,68,72,76,82,86,90,94],
                    'AG':[7,10,12,14,16,20,23,25,27,29,33,36,38,40,42,45,50,54,58,68,72,76,82,86,90]}
    nonblank_marks = 0
    checkmark_count = 0
    checkmark_counts_by_card = {}
    for sheet in wb:
        if sheet.title == 'summary':
            continue
        sheet_checkmark_count = 0
        for column, selected_rows in scoring_rows.items():
            for row in selected_rows:
                val = sheet[f'{column}{row}'].value
                # 模板有两种行：六次ISBLANK的等级选择，以及仅C/S一格的项目选择。
                # 不能把所有行一概当七格；例如E10是项目标签，不是勾选格。
                refs = re.findall(r'ISBLANK\(\s*\$?([A-Z]+)\$?(\d+)\s*\)',str(val),flags=re.IGNORECASE)
                checks = {f'{c.upper()}{r}' for c,r in refs if int(r)==row}
                if len(checks)>1:
                    # 等级公式把未选正等级默认记为0，因此0级C/S格未显式引用。
                    checks.add(f'{"C" if column=="Q" else "S"}{row}')
                nonblank_marks += sum(sheet[address].value is not None for address in checks)
                checkmark_count += len(checks)
                sheet_checkmark_count += len(checks)
                if val is not None:
                    card_rows.append({'sheet':sheet.title,'object':sheet['S3'].value,'cell':f'{column}{row}',
                                      'formula_or_value':val,'cached_value':cached[sheet.title][f'{column}{row}'].value,
                                      'check_cells':';'.join(sorted(checks,key=lambda address:sheet[address].column)),
                                      'n_check_cells':len(checks),
                                      'nonblank_check_cells':sum(sheet[address].value is not None for address in checks)})
        checkmark_counts_by_card[sheet.title] = sheet_checkmark_count
    pd.DataFrame(card_rows).to_csv(OUT / 'data_audit_template_cells.csv', index=False, encoding='utf-8-sig')
    summary['D03_sheet_names'] = wb.sheetnames
    summary['D03_cards'] = len(wb.sheetnames)-1
    summary['D03_checkmark_cells'] = checkmark_count
    summary['D03_checkmark_cells_per_card'] = checkmark_counts_by_card
    summary['D03_nonblank_checkmarks'] = nonblank_marks
    summary['critical_scale'] = json.loads(transforms[transforms.var_id.isin(['V004','V039','V021','V109'])].to_json(orient='records'))
    critical = x[['object_id','regime','V004','V039','V021','V109']].copy()
    critical['V004'] = scale['V004']
    critical['V021'] = scale['V021']
    critical.to_csv(OUT / 'data_audit_critical_native_observed.csv', encoding='utf-8-sig')
    refdata = sheets['ref_attribution'].query("kind == 'data'").set_index('object')
    summary['ref_attribution_data_D01_max_abs_error'] = float(np.nanmax(np.abs(refdata[features]-xnum)))
    summary['ref_attribution_data_missing_cells'] = int(refdata[features].isna().sum().sum())
    # 假设性相加：附件没有说明归因属于哪个目标、哪种输出刻度。
    # 因而这里只追查相加后的数能否作为标度线索，不称为已核验的模型预测。
    attribution = sheets['ref_attribution'].query("kind == 'attribution'").set_index('object')
    additive = pd.DataFrame(index=attribution.index)
    additive['attribution_sum'] = attribution[features].sum(axis=1,min_count=len(features))
    additive['base_prediction'] = attribution['base_prediction']
    additive['hypothetical_scalar'] = additive['attribution_sum'] + additive['base_prediction']
    additive['distance_to_nearest_integer'] = (additive['hypothetical_scalar']-additive['hypothetical_scalar'].round()).abs()
    additive.to_csv(OUT / 'data_audit_attribution_hypothetical_scalar.csv',encoding='utf-8-sig')
    alignments = []
    for target in targets:
        paired = pd.concat([additive['hypothetical_scalar'],ynum[target]],axis=1).dropna()
        scalar, observed = paired['hypothetical_scalar'].to_numpy(), paired[target].to_numpy()
        # 两参数仿射拟合只描述同一批观测的相容程度，不是留出预测精度。
        design = np.column_stack([np.ones(len(paired)),scalar])
        intercept, slope = np.linalg.lstsq(design,observed,rcond=None)[0]
        residual = observed - design @ np.array([intercept,slope])
        mse = float(np.mean(residual**2))
        variance = float(np.var(observed))
        alignments.append({'target':target,'n_pairs':len(paired),
                           'pearson_r':float(paired.corr(method='pearson').iloc[0,1]),
                           'spearman_r':float(paired.corr(method='spearman').iloc[0,1]),
                           'affine_intercept':float(intercept),'affine_slope':float(slope),
                           'affine_RMSE_D04_units':float(np.sqrt(mse)),
                           'affine_max_abs_residual_D04_units':float(np.max(np.abs(residual))),
                           'affine_R2_in_sample_descriptive':float(1-mse/variance) if variance>0 else np.nan,
                           'affine_RMSE_divided_by_target_sd':float(np.sqrt(mse/variance)) if variance>0 else np.nan})
    alignment_frame = pd.DataFrame(alignments)
    alignment_frame.to_csv(OUT / 'data_audit_attribution_target_alignment.csv',index=False,encoding='utf-8-sig')
    strongest = alignment_frame.loc[alignment_frame['pearson_r'].abs().idxmax()]
    summary['attribution_hypothetical_additive_scalar'] = {
        'scope':'hypothetical additive reconstruction; target and output scale undocumented',
        'n_objects':len(additive),'n_complete_sums':int(additive['hypothetical_scalar'].notna().sum()),
        'base_prediction_nunique':int(additive['base_prediction'].nunique()),
        'base_prediction_min':float(additive['base_prediction'].min()),
        'base_prediction_max':float(additive['base_prediction'].max()),
        'candidate_min':float(additive['hypothetical_scalar'].min()),
        'candidate_max':float(additive['hypothetical_scalar'].max()),
        'n_near_integer_abs_error_le_1e_6':int(additive['distance_to_nearest_integer'].le(1e-6).sum()),
        'strongest_absolute_pearson_target':strongest['target'],
        'strongest_absolute_pearson_r':float(strongest['pearson_r']),
        'strongest_target_affine_RMSE_D04_units':float(strongest['affine_RMSE_D04_units']),
        'strongest_target_affine_max_abs_residual_D04_units':float(strongest['affine_max_abs_residual_D04_units']),
        'strongest_target_affine_R2_in_sample_descriptive':float(strongest['affine_R2_in_sample_descriptive']),
    }
    with open(OUT / 'data_audit_summary.json', 'w', encoding='utf-8') as f:
        json.dump(summary, f, ensure_ascii=False, indent=2)
    figures = OUT.parent / 'figures'
    figures.mkdir(exist_ok=True)
    fig, axes = plt.subplots(1,3,figsize=(14,4.1),layout='constrained')
    v='V030'
    b=scale[v].dropna()
    grid=np.linspace(b.min(),b.max(),200)
    axes[0].scatter(b, xnum.loc[b.index,v], s=12, alpha=.4, label='paired observations')
    axes[0].plot(grid,np.log10(grid),label='log10(native)',color='black')
    axes[0].plot(grid,np.log(grid),label='ln(native)',color='tomato',linestyle='--')
    axes[0].set(xlabel='Reference native V030',ylabel='D01 V030',title='Paired scale identifies the base')
    axes[0].legend(fontsize=8)
    b=scale['V021']
    axes[1].scatter(b,xnum['V021'],s=14,alpha=.5)
    axes[1].set(xlabel='Reference native V021',ylabel='D01 V021',title='168 native zeros share one finite floor')
    axes[1].axhline(xnum.loc[b.eq(0),'V021'].iloc[0],linestyle=':',color='tomato')
    axes[2].scatter(scale['V095'],xnum['V095'],s=16,alpha=.6)
    lim=max(scale['V095'].max(),xnum['V095'].max())
    axes[2].plot([0,lim],[0,lim],color='black',linestyle='--')
    axes[2].set(xlabel='Reference native V095',ylabel='D01 V095',title='Same scale does not mean identical data')
    fig.savefig(figures / 'data_audit_scale_evidence.png',dpi=180)
    plt.close(fig)
    print(json.dumps(summary, ensure_ascii=False, indent=2))
    print(pd.DataFrame(corr_rows).to_string(index=False))
    print('Per-column scale details saved in data_audit_scale_candidates.csv')


if __name__ == '__main__':
    main()
