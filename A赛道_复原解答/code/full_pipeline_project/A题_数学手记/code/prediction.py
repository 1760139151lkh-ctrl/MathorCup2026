"""问题1：先预测可观测的连续评分；工作等级另附假设，绝不冒称原始评分。

从项目根目录运行：python A题_数学手记/code/prediction.py
所有学习到的填补值、裁剪边界、均值和标准差都只来自当前训练折。
"""
from pathlib import Path
import json
import math
import sys
from functools import lru_cache
from copy import copy
import numpy as np
import pandas as pd
import joblib
from sklearn.base import BaseEstimator, TransformerMixin
from sklearn.pipeline import Pipeline
from sklearn.compose import TransformedTargetRegressor
from sklearn.preprocessing import StandardScaler
from sklearn.linear_model import Ridge
from sklearn.cross_decomposition import PLSRegression
from sklearn.model_selection import KFold, GroupKFold, train_test_split
from sklearn.metrics import r2_score, mean_absolute_error
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from scipy.stats import beta
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parents[1]
RESULTS = WORK / "results"
FIGURES = WORK / "figures"
FEATURES = [f"V{i:03d}" for i in range(1, 232)]
TARGETS = [f"Y{i:02d}" for i in range(1, 51)]
SEED = 20261026
CANDIDATES = [("ridge", a) for a in (0.1, 1., 10., 100., 1000.)]
CANDIDATES += [("pls", k) for k in (2, 5, 10, 20)]


def feature_array(X):
    """外部调用仍给原D01尺度；列顺序在这里统一，元数据不进入模型。"""
    if isinstance(X, pd.DataFrame):
        missing = set(FEATURES) - set(X.columns)
        if missing:
            raise ValueError(f"缺少特征列：{sorted(missing)}")
        a = X.loc[:, FEATURES].to_numpy(dtype=float)
    else:
        a = np.asarray(X, dtype=float)
    if a.ndim != 2 or a.shape[1] != len(FEATURES):
        raise ValueError("输入必须是n×231原D01尺度的二维特征矩阵")
    return np.where(np.isfinite(a), a, np.nan)


class RobustPreprocessor(TransformerMixin, BaseEstimator):
    """先补训练中位数，再裁剪训练1%/99%分位，最后换成标准差刻度。

    裁剪不是宣布极端值错误，而是限制少数巨大数值对线性模型的支配。
    原始数据不改写；生成样本的物理阈值也不在这个变换空间里检查。
    """
    def fit(self, X, y=None):
        a = feature_array(X)
        self.median_ = np.nanmedian(a, axis=0)
        self.median_ = np.where(np.isfinite(self.median_), self.median_, 0.)
        filled = np.where(np.isnan(a), self.median_, a)
        self.low_, self.high_ = np.quantile(filled, [0.01, 0.99], axis=0)
        clipped = np.clip(filled, self.low_, self.high_)
        self.mean_ = clipped.mean(axis=0)
        sd = clipped.std(axis=0)
        self.scale_ = np.where(sd > 1e-12, sd, 1.)
        self.n_features_in_ = a.shape[1]
        return self

    def transform(self, X):
        a = feature_array(X)
        filled = np.where(np.isnan(a), self.median_, a)
        return (np.clip(filled, self.low_, self.high_) - self.mean_) / self.scale_


def make_estimator(kind, value):
    reg = Ridge(alpha=value) if kind == "ridge" else PLSRegression(n_components=int(value), scale=False, max_iter=1000)
    pipe = Pipeline([("preprocess", RobustPreprocessor()), ("reg", reg)])
    # 50个目标的训练尺度不同。让每个目标都以自己的训练标准差衡量，
    # 避免原来数值波动较大的目标独占PLS寻找共同方向的过程。
    return TransformedTargetRegressor(regressor=pipe, transformer=StandardScaler())


def load_xy():
    x = pd.read_csv(ROOT / "data/D01.csv")
    y = pd.read_csv(ROOT / "data/D04.csv")
    # object是字典约定的唯一序号；object_id在两表之间全部冲突，不能作为连接键。
    if x.object.duplicated().any() or y.object.duplicated().any():
        raise ValueError("object键不唯一，停止连接")
    aligned = y.set_index("object").reindex(x.object)
    if aligned[TARGETS].isna().any().any():
        raise ValueError("目标缺失或object无法对齐，需重新调查")
    if not np.array_equal(x.regime.to_numpy(), aligned.regime.to_numpy()):
        raise ValueError("相同object的类别冲突，需人工核查")
    return x.reset_index(drop=True), aligned[TARGETS].reset_index(drop=True)


def tune(X, Y, groups=None, seed=SEED):
    splitter = KFold(3, shuffle=True, random_state=seed) if groups is None else GroupKFold(3)
    splits = list(splitter.split(X, Y, groups))
    records = []
    for kind, value in CANDIDATES:
        losses = []
        for tr, va in splits:
            model = make_estimator(kind, value).fit(X[tr], Y[tr])
            pred = model.predict(X[va])
            scale = np.maximum(Y[tr].std(axis=0), 1e-12)
            losses.append(float(np.mean(((Y[va] - pred) / scale) ** 2)))
        records.append({"family": kind, "parameter": float(value), "inner_normalized_mse": float(np.mean(losses))})
    table = pd.DataFrame(records).sort_values(["inner_normalized_mse", "family", "parameter"])
    best = table.iloc[0]
    return (best.family, best.parameter), table


class PredictionBundle:
    """把预测器和工作等级的约定绑在一起，生成模块无需重新定义刻度。"""
    def __init__(self, estimator, Y, selection):
        self.estimator = estimator
        self.selection_ = selection
        self.feature_names_ = FEATURES.copy()
        self.target_names_ = TARGETS.copy()
        self.grade_location_ = np.asarray(Y).mean(axis=0)
        self.grade_scale_ = np.maximum(np.asarray(Y).std(axis=0), 1e-12)
        reg = estimator.regressor_.named_steps["reg"]
        coef = np.asarray(reg.coef_)
        if coef.shape == (231, 50):
            coef = coef.T
        intercept = np.asarray(reg.intercept_).reshape(-1)
        if isinstance(reg, PLSRegression):
            intercept = intercept - coef @ reg._x_mean
        scale = estimator.transformer_.scale_
        self.linear_coef_ = coef * scale[:, None]
        self.linear_intercept_ = estimator.transformer_.mean_ + intercept * scale

    def predict(self, X):
        return self.estimator.predict(feature_array(X))


def fit_final_model(X=None, Y=None, random_state=SEED):
    if X is None or Y is None:
        loaded_x, loaded_y = load_xy()
        X, Y = loaded_x[FEATURES], loaded_y[TARGETS]
    a, b = feature_array(X), np.asarray(Y, dtype=float)
    selection, table = tune(a, b, seed=random_state)
    estimator = make_estimator(*selection).fit(a, b)
    bundle = PredictionBundle(estimator, b, selection)
    bundle.tuning_ = table
    return bundle


@lru_cache(maxsize=1)
def load_final_model():
    path = RESULTS / "question1_model.joblib"
    if not path.exists():
        raise FileNotFoundError("请先运行prediction.py生成question1_model.joblib")
    return joblib.load(path)


def predict_continuous(X, model=None):
    bundle = load_final_model() if model is None else model
    return pd.DataFrame(bundle.predict(X), columns=TARGETS, index=X.index if isinstance(X, pd.DataFrame) else None)


def working_grades(scores, model=None, slope=1.):
    """人为约定：D04均值映到3，每增加一个训练标准差增加slope级。

    这不是原始评分的反演；改变约定可以改变等级，尤其会改变Y38≥2判断。
    floor(t+0.5)明确了半级向上取整，避免numpy银行家舍入的隐含约定。
    """
    bundle = load_final_model() if model is None else model
    a = scores[TARGETS].to_numpy() if isinstance(scores, pd.DataFrame) else np.asarray(scores)
    raw = 3. + slope * (a - bundle.grade_location_) / bundle.grade_scale_
    grades = np.clip(np.floor(raw + 0.5), 0, 6).astype(int)
    return pd.DataFrame(grades, columns=TARGETS, index=scores.index if isinstance(scores, pd.DataFrame) else None)


def metric_rows(Y, pred, model_name, protocol):
    rows = []
    for j, target in enumerate(TARGETS):
        residual = Y[:, j] - pred[:, j]
        rows.append({"target": target, "model": model_name, "protocol": protocol,
                     "mae": float(np.mean(np.abs(residual))),
                     "rmse": float(np.sqrt(np.mean(residual ** 2))),
                     "r2": float(r2_score(Y[:, j], pred[:, j]))})
    return rows


def nested_cv(X, Y, groups=None, protocol_name=None):
    outer = KFold(5, shuffle=True, random_state=SEED) if groups is None else GroupKFold(5)
    preds = {name: np.empty_like(Y) for name in ("mean", "ridge", "pls", "selected")}
    choices, fold_labels = [], np.empty(len(Y), dtype=int)
    oof_grades = np.empty_like(Y, dtype=int)
    for fold, (tr, te) in enumerate(outer.split(X, Y, groups), start=1):
        selection, inner = tune(X[tr], Y[tr], None if groups is None else groups[tr], seed=SEED + fold)
        fold_labels[te] = fold
        preds["mean"][te] = Y[tr].mean(axis=0)
        for family in ("ridge", "pls"):
            best = inner[inner.family == family].iloc[0]
            fitted = make_estimator(family, best.parameter).fit(X[tr], Y[tr])
            preds[family][te] = fitted.predict(X[te])
        preds["selected"][te] = preds[selection[0]][te]
        grade_value = 3. + (preds["selected"][te] - Y[tr].mean(axis=0)) / np.maximum(Y[tr].std(axis=0),1e-12)
        oof_grades[te] = np.clip(np.floor(grade_value+0.5),0,6).astype(int)
        inner["outer_fold"] = fold
        inner["selected_in_fold"] = ((inner.family == selection[0]) & (inner.parameter == selection[1]))
        choices.append(inner)
        label = protocol_name or ("random" if groups is None else "regime_holdout")
        print(f"{label} fold {fold}: {selection}", flush=True)
    return preds, fold_labels, pd.concat(choices, ignore_index=True), oof_grades


def conformal(X, Y, objects):
    all_idx = np.arange(len(Y))
    tr, held = train_test_split(all_idx, test_size=100, random_state=SEED + 99)
    ca, te = train_test_split(held, test_size=50, random_state=SEED + 100)
    selection, _ = tune(X[tr], Y[tr], seed=SEED + 101)
    model = make_estimator(*selection).fit(X[tr], Y[tr])
    residuals = np.abs(Y[ca] - model.predict(X[ca]))
    rank = math.ceil((len(ca) + 1) * 0.9)
    q = np.sort(residuals, axis=0)[rank - 1]
    pred = model.predict(X[te])
    covered = np.abs(Y[te] - pred) <= q
    records = []
    for j, target in enumerate(TARGETS):
        count, n = int(covered[:, j].sum()), len(te)
        lo = 0. if count == 0 else float(beta.ppf(0.025, count, n-count+1))
        hi = 1. if count == n else float(beta.ppf(0.975, count+1, n-count))
        records.append({"target": target, "calibration_n": len(ca), "test_n": n,
                        "nominal_marginal_coverage": 0.9, "test_covered": count,
                        "test_coverage": count/n, "coverage_ci95_low": lo,
                        "coverage_ci95_high": hi, "interval_width": float(2*q[j]),
                        "calibration_order_statistic": rank})
    long = []
    for i, obj in enumerate(objects[te]):
        for j, target in enumerate(TARGETS):
            long.append({"object": int(obj), "target": target, "observed_D04": float(Y[te[i], j]),
                         "prediction": float(pred[i, j]), "lower": float(pred[i, j]-q[j]),
                         "upper": float(pred[i, j]+q[j]), "covered": bool(covered[i, j])})
    pd.DataFrame(records).to_csv(RESULTS / "question1_conformal_metrics.csv", index=False)
    pd.DataFrame(long).to_csv(RESULTS / "question1_conformal_test_intervals.csv", index=False)
    train_scale = np.maximum(Y[tr].std(axis=0), 1e-12)
    joint_scores = np.max(residuals / train_scale, axis=1)
    joint_q = float(np.sort(joint_scores)[rank-1])
    joint_halfwidth = joint_q * train_scale
    joint_covered = np.all(np.abs(Y[te]-pred) <= joint_halfwidth, axis=1)
    pd.DataFrame({"target":TARGETS,"joint_box_halfwidth":joint_halfwidth,"calibration_max_residual_quantile":joint_q,"train_target_scale":train_scale}).to_csv(RESULTS/"question1_conformal_joint_widths.csv",index=False)
    joint_rows=[]
    for i,obj in enumerate(objects[te]):
        joint_rows.append({"object":int(obj),"all_50_covered_by_joint_box":bool(joint_covered[i]),"all_50_covered_by_marginal_intervals":bool(np.all(covered[i])),"nominal_joint_coverage":0.9})
    pd.DataFrame(joint_rows).to_csv(RESULTS/"question1_conformal_joint_coverage.csv",index=False)
    pd.DataFrame({"object":objects, "split":np.select([np.isin(all_idx,tr),np.isin(all_idx,ca)],["train","calibration"],default="test")}).to_csv(RESULTS / "question1_conformal_split.csv",index=False)
    joint_count = int(joint_covered.sum())
    joint_ci = [float(beta.ppf(.025,joint_count,len(te)-joint_count+1)),float(beta.ppf(.975,joint_count+1,len(te)-joint_count))]
    return {"mean_marginal_coverage":float(covered.mean()),"joint_box_covered":joint_count,
            "joint_box_test_n":len(te),"joint_box_test_coverage":float(joint_covered.mean()),
            "joint_box_coverage_ci95":joint_ci,"all50_covered_by_marginal_intervals":float(np.all(covered,axis=1).mean())}, selection


def card_mappings():
    """字典第3节逐项对应；None表示模板有项目、D04无目标，不能擅自补分。"""
    rows_a = [7,10,12,14,16,20,23,25,27,29,33,36,38,40,42,45,50,54,58,62,68,72,76,82,86,90,94]
    ys_a = [1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,None,48,32,49,47,36,37,34,40,41,42,50]
    rows_f = [7,10,12,14,16,20,23,25,27,29,33,36,38,40,42,45,50,54,58,68,72,76,82,86,90]
    ys_f = [16,17,18,19,20,21,22,23,24,25,26,27,28,29,30,None,46,33,31,39,35,38,43,44,45]
    mapping={}
    for col,rows,targets in (("Q",rows_a,ys_a),("AG",rows_f,ys_f)):
        mapping.update({f"{col}{row}":None if target is None else f"Y{target:02d}" for row,target in zip(rows,targets)})
    return mapping


def build_card_payload(grades,objects):
    """数据载荷供artifact-tool保留原模板排版后填入，也供openpyxl独立复现。"""
    wb = openpyxl.load_workbook(ROOT / "data/D03.xlsx")
    score_by_object = grades.copy()
    score_by_object.index = objects
    note = "PROTOCOL A: clip(floor(3+(z-mu)/sd+0.5),0,6); full D04 mu/sd. A-16/F-16 unavailable."
    short_note = "PROTOCOL A: relative grades; original scale unidentified."
    mapping = card_mappings()
    cards=[]
    for ws in wb.worksheets:
        if not ws.title.startswith("object_"):
            continue
        obj = int(ws["S3"].value)
        cells={"S3":obj,"B2":short_note,"B104":note}
        cells.update({address:"N/A" if target is None else int(score_by_object.loc[obj,target]) for address,target in mapping.items()})
        cards.append({"sheet":ws.title,"object":obj,"cells":cells})
    summary = wb["summary"]
    summary_updates=[]
    for row in range(2, summary.max_row+1):
        obj = summary.cell(row,1).value
        if obj in score_by_object.index:
            values=[int(obj)]+[int(score_by_object.loc[obj,target]) for target in TARGETS]+[note]
            summary_updates.append({"range":f"A{row}:AZ{row}","values":[values]})
    # 当前模板summary只有第1-13行；第15行原本为空。
    summary_updates.append({"range":"A15","values":[[short_note]]})
    return {"cards":cards,"summary":summary_updates,"note":note,"short_note":short_note}


def write_cards(grades,objects):
    wb=openpyxl.load_workbook(ROOT/"data/D03.xlsx")
    payload=build_card_payload(grades,objects)
    for card in payload["cards"]:
        for address,value in card["cells"].items():
            wb[card["sheet"]][address]=value
        ws=wb[card["sheet"]]
        for address,row in (("B2",2),("B104",104)):
            font=copy(ws[address].font)
            font.sz=10
            ws[address].font=font
            ws.row_dimensions[row].height=20
    for update in payload["summary"]:
        top_left=update["range"].split(":")[0]
        cell=wb["summary"][top_left]
        for i,row in enumerate(update["values"]):
            for j,value in enumerate(row):
                wb["summary"].cell(cell.row+i,cell.column+j,value)
    font=copy(wb["summary"]["A15"].font)
    font.sz=10
    wb["summary"]["A15"].font=font
    wb["summary"].row_dimensions[15].height=20
    wb["summary"].column_dimensions["A"].width=12
    for col in range(2,52):
        wb["summary"].column_dimensions[openpyxl.utils.get_column_letter(col)].width=6
    wb["summary"].column_dimensions["AZ"].width=120
    wb["summary"].freeze_panes="B2"
    wb.save(RESULTS / "question1_working_grades_D03_ASSUMPTION.xlsx")


def cards_main():
    """只重建评分卡：python prediction.py --cards-only；先运行主计算。"""
    x,_=load_xy()
    bundle=load_final_model()
    write_cards(working_grades(predict_continuous(x,bundle),bundle),x.object.to_numpy())


def main():
    RESULTS.mkdir(parents=True,exist_ok=True)
    FIGURES.mkdir(parents=True,exist_ok=True)
    x, y = load_xy()
    X, Y = feature_array(x), y.to_numpy()
    objects = x.object.to_numpy()
    random_preds, folds, random_choices, oof_grades = nested_cv(X,Y)
    group_preds, group_folds, group_choices, _ = nested_cv(X,Y,x.regime.to_numpy())
    d04 = pd.read_csv(ROOT/"data/D04.csv").set_index("object").reindex(objects)
    id_preds, id_folds, id_choices, _ = nested_cv(X,Y,d04.object_id.to_numpy(),protocol_name="D04_ID_holdout_sensitivity")
    metrics = []
    for protocol, preds in (("nested_random_5x3",random_preds),("nested_regime_holdout_5x3",group_preds),("nested_D04_ID_holdout_5x3_sensitivity",id_preds)):
        for name,pred in preds.items():
            metrics.extend(metric_rows(Y,pred,name,protocol))
    metrics = pd.DataFrame(metrics)
    baseline = metrics[metrics.model == "mean"].set_index(["protocol","target"]).mae
    metrics["mae_improvement_over_mean"] = [1-r.mae/baseline.loc[(r.protocol,r.target)] for r in metrics.itertuples()]
    metrics.to_csv(RESULTS / "question1_metrics.csv",index=False)
    random_choices["protocol"] = "random"
    group_choices["protocol"] = "regime_holdout"
    id_choices["protocol"] = "D04_ID_holdout_sensitivity"
    pd.concat([random_choices,group_choices,id_choices]).to_csv(RESULTS / "question1_inner_selection.csv",index=False)
    pd.DataFrame({"object":objects,"random_outer_fold":folds,"regime_outer_fold":group_folds,"D04_ID_outer_fold":id_folds}).to_csv(RESULTS / "question1_cv_folds.csv",index=False)
    oof = pd.DataFrame(random_preds["selected"],columns=TARGETS)
    pd.concat([x.iloc[:,:3],oof],axis=1).to_csv(RESULTS / "question1_continuous_oof.csv",index=False)
    pd.concat([x.iloc[:,:3],pd.DataFrame(oof_grades,columns=TARGETS)],axis=1).to_csv(RESULTS / "question1_working_grades_oof_ASSUMPTION.csv",index=False)
    pd.concat([x.iloc[:,:3],pd.DataFrame(id_preds["selected"],columns=TARGETS)],axis=1).to_csv(RESULTS / "question1_continuous_D04_ID_oof_sensitivity.csv",index=False)
    bundle = fit_final_model(X,Y)
    joblib.dump(bundle,RESULTS / "question1_model.joblib")
    load_final_model.cache_clear()
    bundle.tuning_.to_csv(RESULTS / "question1_final_selection.csv",index=False)
    final = predict_continuous(x,bundle)
    pd.concat([x.iloc[:,:3],final],axis=1).to_csv(RESULTS / "question1_continuous_fitted.csv",index=False)
    grades = working_grades(final,bundle)
    pd.concat([x.iloc[:,:3],grades],axis=1).to_csv(RESULTS / "question1_working_grades_ASSUMPTION.csv",index=False)
    # 用观测D04应用同一人为规则，供区分“有数据对象的描述”与“模型预测”。
    observed_grades = working_grades(y,bundle)
    pd.concat([x.iloc[:,:3],observed_grades],axis=1).to_csv(RESULTS / "question1_observed_working_grades_ASSUMPTION.csv",index=False)
    pd.DataFrame({"target":TARGETS,"location_D04":bundle.grade_location_,"scale_D04":bundle.grade_scale_,"midpoint_grade":3,"primary_slope":1.,"status":"ASSUMPTION: not original scale recovery"}).to_csv(RESULTS / "question1_grade_rule_ASSUMPTION.csv",index=False)
    sensitivity = []
    for slope in (0.5,1.,1.5,2.):
        variant = working_grades(final,bundle,slope=slope)
        for j,target in enumerate(TARGETS):
            sensitivity.append({"slope":slope,"target":target,"fraction_changed_from_slope1":float((variant[target]!=grades[target]).mean()),"mean_working_grade":float(variant[target].mean()),"K_threshold_status_changed_fraction":float(((variant.Y38>=2)!=(grades.Y38>=2)).mean()) if target=="Y38" else np.nan})
    pd.DataFrame(sensitivity).to_csv(RESULTS / "question1_grade_sensitivity_ASSUMPTION.csv",index=False)
    # D03模板填卡由主流程在可视化检查后完成；write_cards保留独立Python复现入口。
    cp_metrics, cp_selection = conformal(X,Y,objects)
    summary = {"alignment":"object; object_id conflicts in every row; regime agrees", "n":len(Y),"features":231,"targets":50,
               "D03_cards":12,"D03_has_observed_grades":False,"raw_grade_identified":False,
               "final_selection":{"family":bundle.selection_[0],"parameter":float(bundle.selection_[1])},
               "conformal":cp_metrics,"conformal_selection":{"family":cp_selection[0],"parameter":float(cp_selection[1])},
               "grade_warning":"All integer grades are an explicit working convention, not recovered original labels. No original-grade accuracy is estimable.",
               "conformal_warning":"Separate 90% marginal intervals and 90% joint max-residual box apply to exchangeable D04 target rows; no class-conditional, generated-OOD or causal safety guarantee. Repeated D04 IDs have unknown semantics, so entity dependence is unresolved."}
    (RESULTS / "question1_summary.json").write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding="utf-8")
    selected = metrics[(metrics.model=="selected")&(metrics.protocol=="nested_random_5x3")].set_index("target")
    grouped = metrics[(metrics.model=="selected")&(metrics.protocol=="nested_regime_holdout_5x3")].set_index("target")
    plt.rcParams.update({"font.size":10,"figure.dpi":150})
    fig,ax=plt.subplots(figsize=(14,4.5))
    axis=np.arange(50)
    ax.bar(axis-0.18,selected.loc[TARGETS,"r2"],width=.36,label="Random nested CV")
    ax.bar(axis+.18,grouped.loc[TARGETS,"r2"],width=.36,label="Regime-held-out nested CV",alpha=.8)
    ax.axhline(0,color="black",lw=.7);ax.set_xticks(axis);ax.set_xticklabels(TARGETS,rotation=90)
    ax.set_ylabel("Out-of-fold R-squared");ax.legend();fig.tight_layout();fig.savefig(FIGURES/"question1_cv_r2.png");plt.close(fig)
    fig,axes=plt.subplots(1,3,figsize=(12,3.8))
    for ax,target in zip(axes,["Y01","Y38","Y50"]):
        j=TARGETS.index(target);ax.scatter(Y[:,j],oof[target],s=12,alpha=.65)
        lim=[min(Y[:,j].min(),oof[target].min()),max(Y[:,j].max(),oof[target].max())]
        ax.plot(lim,lim,"k--",lw=1);ax.set(xlabel="Observed D04 continuous value",ylabel="Out-of-fold prediction",title=target)
    fig.tight_layout();fig.savefig(FIGURES/"question1_oof_scatter.png");plt.close(fig)
    fig,ax=plt.subplots(figsize=(8,4))
    sens=pd.DataFrame(sensitivity).groupby("slope").fraction_changed_from_slope1.mean()
    ax.plot(sens.index,sens.values,"o-");ax.set(xlabel="Arbitrary grade-map slope",ylabel="Fraction of 12,500 model grades changed",ylim=(0,1))
    ax.set_title("Sensitivity of the working convention, not original-grade accuracy")
    fig.tight_layout();fig.savefig(FIGURES/"question1_grade_sensitivity.png");plt.close(fig)
    print(json.dumps(summary,ensure_ascii=False,indent=2),flush=True)
    print(metrics.groupby(["protocol","model"])[["mae","r2","mae_improvement_over_mean"]].mean().to_string(),flush=True)


if __name__ == "__main__":
    # 让joblib中的自定义类来自可重新import的prediction模块，而非临时__main__。
    import prediction as canonical
    if "--cards-only" in sys.argv:
        canonical.cards_main()
    else:
        canonical.main()
